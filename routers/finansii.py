from fastapi import APIRouter, Request, Form, Depends, BackgroundTasks, Body, HTTPException
from fastapi.responses import (
    Response, RedirectResponse, JSONResponse, FileResponse
)
from fastapi.templating import Jinja2Templates
import openpyxl
from openpyxl.utils import get_column_letter
from datetime import datetime, date
from calendar import monthrange
import os
import re
import unicodedata
import asyncio
import uuid
import time
import json

from db_ifx import informix_cursor
from auth.role_utils import has_any_role

router = APIRouter()

# ---------------------------
# Templates
# ---------------------------
templates = Jinja2Templates(directory=os.path.join(os.path.dirname(__file__), "..", "templates"))

EXPORT_DIR = "/opt/siglife-reporting/exports/finansii"
JOB_DIR = "/opt/siglife-reporting/exports/finansii/jobs"
FINANCIAL_TEMPLATES_PATH = os.path.join(
    os.path.dirname(__file__),
    "..",
    "data",
    "financial_statement_templates.json"
)
FINANCIAL_DEFINITIONS_PATH = os.path.join(
    os.path.dirname(__file__),
    "..",
    "data",
    "financial_statement_definitions.json"
)

# ---------------------------
# Session check
# ---------------------------
def get_current_user(request: Request):
    user = request.session.get("user")
    if not user:
        return RedirectResponse(url="/siglife-report/login", status_code=303)
    return user

def get_finansii_user(request: Request):
    user = get_current_user(request)
    if isinstance(user, RedirectResponse):
        return user
    if not has_any_role(user, "finance", "finansii"):
        raise HTTPException(status_code=403, detail="Access denied")
    return user

def get_finansiski_izvestai_user(request: Request):
    user = get_current_user(request)
    if isinstance(user, RedirectResponse):
        return user
    if not has_any_role(user, "finance", "finansiski_izvestai"):
        raise HTTPException(status_code=403, detail="Access denied")
    return user

def safe_ascii_filename(name: str) -> str:
    name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
    return name or "report"

def _normalize_account_masks(raw_value: str) -> list[str]:
    masks = []
    normalized_value = str(raw_value or "")
    normalized_value = normalized_value.replace(" и ", ",")
    for part in normalized_value.split(","):
        mask = part.strip().replace("*", "%").replace(" ", "")
        mask = re.sub(r"[дД]$", "", mask)
        if not mask:
            continue
        if not re.fullmatch(r"[0-9%]+", mask):
            raise ValueError(f"Invalid account mask: {part}")
        if "%" not in mask:
            mask = f"{mask}%"
        masks.append(mask)
    return masks

def _amount_for_type(debit: float, credit: float, definition_type: str) -> float:
    if definition_type in ("Расход", "Одлив", "Актива"):
        return debit - credit
    return credit - debit

def _fetch_statement_definition_amount(cursor, year: int, month_from: int, month_to: int, masks: list[str]):
    if not masks:
        return 0.0, 0.0

    last_day = monthrange(year, month_to)[1]
    like_clause = " OR ".join([f"konto LIKE '{mask}'" for mask in masks])
    sql = f"""
SELECT
    NVL(SUM(iznos_d), 0) AS debit,
    NVL(SUM(iznos_p), 0) AS credit
FROM stavka
WHERE godina = {year}
  AND dat_nalog >= MDY({month_from}, 1, {year})
  AND dat_nalog <= MDY({month_to}, {last_day}, {year})
  AND ({like_clause})
"""
    cursor.execute(sql)
    row = cursor.fetchone()

    if not row:
        return 0.0, 0.0
    return float(row[0] or 0), float(row[1] or 0)

# ---------------------------
# Job JSON store on disk
# ---------------------------
def _job_path(job_id: str) -> str:
    os.makedirs(JOB_DIR, exist_ok=True)
    return os.path.join(JOB_DIR, f"{job_id}.json")

def _write_job(job_id: str, payload: dict):
    p = _job_path(job_id)
    tmp = p + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False)
    os.replace(tmp, p)

def _read_job(job_id: str):
    p = _job_path(job_id)
    if not os.path.exists(p):
        return None
    with open(p, "r", encoding="utf-8") as f:
        return json.load(f)

def _read_financial_definitions():
    if not os.path.exists(FINANCIAL_DEFINITIONS_PATH):
        return []
    with open(FINANCIAL_DEFINITIONS_PATH, "r", encoding="utf-8") as f:
        payload = json.load(f)
    if isinstance(payload, dict):
        return payload.get("definitions") or []
    if isinstance(payload, list):
        return payload
    return []

def _write_financial_definitions(definitions: list[dict]):
    os.makedirs(os.path.dirname(FINANCIAL_DEFINITIONS_PATH), exist_ok=True)
    payload = {
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "definitions": definitions
    }
    tmp_path = FINANCIAL_DEFINITIONS_PATH + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    os.replace(tmp_path, FINANCIAL_DEFINITIONS_PATH)

# ---------------------------
# GET page
# ---------------------------
@router.get("/Finansii")
async def finansii_page(request: Request, user=Depends(get_finansii_user)):
    return templates.TemplateResponse(
        "Finansii.html",
        {
            "request": request,
            "user": user,
            "active": "Finansii",
            "selected_date": date.today().isoformat(),
            "message": ""
        }
    )

@router.get("/FinansiskiIzvestai")
@router.get("/finansiski-izvestai")
@router.get("/finansiski_izvestai")
async def finansiski_izvestai_page(request: Request, user=Depends(get_finansiski_izvestai_user)):
    return templates.TemplateResponse(
        "FinansiskiIzvestai.html",
        {
            "request": request,
            "user": user,
            "active": "FinansiskiIzvestai",
            "current_year": date.today().year,
            "message": ""
        }
    )

@router.post("/FinansiskiIzvestai/preview")
async def finansiski_izvestai_preview(
    payload: dict = Body(...),
    user=Depends(get_finansiski_izvestai_user)
):
    try:
        year = int(payload.get("year"))
        month_from = int(payload.get("month_from"))
        month_to = int(payload.get("month_to"))
        definitions = payload.get("definitions") or _read_financial_definitions()

        if year < 2000 or year > 2100:
            return JSONResponse({"ok": False, "error": "Invalid year"}, status_code=400)
        if month_from < 1 or month_from > 12 or month_to < 1 or month_to > 12:
            return JSONResponse({"ok": False, "error": "Invalid month"}, status_code=400)
        if month_to < month_from:
            return JSONResponse({"ok": False, "error": "Month 'to' cannot be before month 'from'"}, status_code=400)
        if len(definitions) > 1000:
            return JSONResponse({"ok": False, "error": "Too many definitions"}, status_code=400)

        rows = []
        total_debit = 0.0
        total_credit = 0.0
        total_amount = 0.0

        with informix_cursor() as cursor:
            for definition in definitions:
                report = str(definition.get("report") or "").strip()
                code = str(definition.get("code") or "").strip()
                position = str(definition.get("position") or "").strip()
                accounts = str(definition.get("accounts") or "").strip()
                definition_type = str(definition.get("type") or "").strip()

                if not position:
                    continue

                debit = 0.0
                credit = 0.0
                amount = 0.0
                if accounts:
                    masks = _normalize_account_masks(accounts)
                    debit, credit = _fetch_statement_definition_amount(cursor, year, month_from, month_to, masks)
                    amount = _amount_for_type(debit, credit, definition_type)

                total_debit += debit
                total_credit += credit
                total_amount += amount

                rows.append({
                    "report": report,
                    "code": code,
                    "position": position,
                    "accounts": accounts,
                    "type": definition_type,
                    "debit": round(debit, 2),
                    "credit": round(credit, 2),
                    "amount": round(amount, 2)
                })

        return JSONResponse({
            "ok": True,
            "rows": rows,
            "summary": {
                "debit": round(total_debit, 2),
                "credit": round(total_credit, 2),
                "amount": round(total_amount, 2)
            }
        })

    except ValueError as e:
        return JSONResponse({"ok": False, "error": str(e)}, status_code=400)
    except Exception as e:
        return JSONResponse({"ok": False, "error": str(e)}, status_code=500)

@router.get("/FinansiskiIzvestai/templates")
async def finansiski_izvestai_templates(user=Depends(get_finansiski_izvestai_user)):
    if not os.path.exists(FINANCIAL_TEMPLATES_PATH):
        return JSONResponse({"ok": False, "error": "Templates file not found"}, status_code=404)

    with open(FINANCIAL_TEMPLATES_PATH, "r", encoding="utf-8") as f:
        payload = json.load(f)

    return JSONResponse({"ok": True, **payload})

@router.get("/FinansiskiIzvestai/db-status")
async def finansiski_izvestai_db_status(user=Depends(get_finansiski_izvestai_user)):
    try:
        with informix_cursor() as cursor:
            cursor.execute("SELECT FIRST 1 tabname FROM systables")
            row = cursor.fetchone()
        return JSONResponse({
            "ok": True,
            "message": "DB конекцијата е успешна.",
            "sample": str(row[0]) if row else None
        })
    except Exception as e:
        return JSONResponse({
            "ok": False,
            "error": str(e)
        }, status_code=500)

@router.get("/FinansiskiIzvestai/definitions")
async def finansiski_izvestai_get_definitions(user=Depends(get_finansiski_izvestai_user)):
    return JSONResponse({
        "ok": True,
        "definitions": _read_financial_definitions()
    })

@router.post("/FinansiskiIzvestai/definitions")
async def finansiski_izvestai_save_definitions(
    payload: dict = Body(...),
    user=Depends(get_finansiski_izvestai_user)
):
    definitions = payload.get("definitions") or []
    if not isinstance(definitions, list):
        return JSONResponse({"ok": False, "error": "Invalid definitions payload"}, status_code=400)
    if len(definitions) > 1000:
        return JSONResponse({"ok": False, "error": "Too many definitions"}, status_code=400)

    cleaned = []
    for item in definitions:
        if not isinstance(item, dict):
            continue
        cleaned.append({
            "report": str(item.get("report") or "").strip(),
            "code": str(item.get("code") or "").strip(),
            "position": str(item.get("position") or "").strip(),
            "accounts": str(item.get("accounts") or "").strip(),
            "type": str(item.get("type") or "").strip(),
            "debit": item.get("debit"),
            "credit": item.get("credit"),
            "amount": item.get("amount")
        })

    _write_financial_definitions(cleaned)
    return JSONResponse({"ok": True, "saved_count": len(cleaned)})

# ============================================================
# 1) SALDIRANJE (PROCEDURE ONLY)  <-- button "Салдирање"
# Stored as job JSON (disk) to avoid RAM issues
# ============================================================
def _run_saldiranje_procedure(d: date) -> str:
    with informix_cursor() as cursor:
        sql = (
            f"EXECUTE PROCEDURE appuser.sp_saldiranje_dokument("
            f"MDY({d.month},{d.day},{d.year})"
            f");"
        )
        cursor.execute(sql)
        result = cursor.fetchone()  # (status, message)

    status = None
    msg = "OK"
    if result:
        status = int(result[0]) if result[0] is not None else None
        msg = str(result[1]) if len(result) > 1 and result[1] is not None else msg

    # Adjust success codes if needed (commonly 0)
    if status in (0, -1):
        return msg

    raise Exception(f"Грешка при салдирање: {msg} (status={status})")

def _run_job_runner_file(job_id: str, d: date):
    job = _read_job(job_id) or {}
    try:
        _write_job(job_id, {**job, "status": "running"})
        msg = _run_saldiranje_procedure(d)
        job2 = _read_job(job_id) or {}
        _write_job(job_id, {
            **job2,
            "status": "success",
            "message": f"Салдирањето е успешно извршено. ({msg})",
            "finished_at": time.time()
        })
    except Exception as e:
        job2 = _read_job(job_id) or {}
        _write_job(job_id, {
            **job2,
            "status": "error",
            "message": str(e),
            "finished_at": time.time()
        })

@router.post("/Finansii/run-async")
async def finansii_run_async(
    request: Request,
    background_tasks: BackgroundTasks,
    export_date_form: str = Form(...),
    user=Depends(get_finansii_user)
):
    try:
        d = datetime.strptime(export_date_form, "%Y-%m-%d").date()
    except ValueError:
        return JSONResponse({"ok": False, "error": "Invalid date format. Use YYYY-MM-DD."}, status_code=400)

    job_id = str(uuid.uuid4())
    _write_job(job_id, {
        "ok": True,
        "type": "run",
        "status": "queued",
        "message": "",
        "started_at": time.time(),
        "finished_at": None
    })

    background_tasks.add_task(_run_job_runner_file, job_id, d)
    return JSONResponse({"ok": True, "job_id": job_id})

@router.get("/Finansii/job/{job_id}")
async def finansii_run_job_status(job_id: str, user=Depends(get_finansii_user)):
    job = _read_job(job_id)
    if not job:
        return JSONResponse({"ok": False, "error": "Job not found"}, status_code=404)
    return JSONResponse(job)

# ============================================================
# 2) EXPORT (EXCEL ONLY - NO PROCEDURE) <-- button "Салдирање – експорт"
# Also stored as job JSON (disk)
# ============================================================
def _fetch_saldiranje_rows(d: date):
    sql = f"""
SELECT
    konto,
    komitent,
    dokument,
    dok_opis,
    NVL(SUM(iznos_d), 0) AS iznos_d,
    NVL(SUM(iznos_p), 0) AS iznos_p,
    NVL(SUM(dev_iznos_d), 0) AS dev_iznos_d,
    NVL(SUM(dev_iznos_p), 0) AS dev_iznos_p
FROM stavka
WHERE godina = {d.year}
  AND konto LIKE '1200%'
  AND dat_nalog <= MDY({d.month},{d.day},{d.year})
GROUP BY 1,2,3,4
HAVING (
    ABS(NVL(SUM(iznos_d),0) - NVL(SUM(iznos_p),0)) < 5
 OR ABS(NVL(SUM(dev_iznos_d),0) - NVL(SUM(dev_iznos_p),0)) < 1
);
"""
    with informix_cursor() as cursor:
        cursor.execute(sql)
        return cursor.fetchall()

def _build_excel_to_server(rows, export_date: date) -> str:
    columns = [
        "konto", "komitent", "dokument", "dok_opis",
        "iznos_d", "iznos_p", "dev_iznos_d", "dev_iznos_p",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Saldiranje_{export_date.strftime('%Y_%m_%d')}"
    ws.append(columns)

    for r in rows:
        ws.append(list(r))

    for i, col_name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(16, len(col_name) + 2)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=5, max_col=8):
        for cell in row:
            if cell.value is not None:
                cell.number_format = "#,##0.00"

    os.makedirs(EXPORT_DIR, exist_ok=True)
    filename = safe_ascii_filename(f"Saldiranje_{export_date.isoformat()}.xlsx")
    file_path = os.path.join(EXPORT_DIR, filename)

    wb.save(file_path)
    return file_path

def _export_job_runner_file(job_id: str, d: date):
    job = _read_job(job_id) or {}
    try:
        _write_job(job_id, {**job, "status": "running"})

        rows = _fetch_saldiranje_rows(d)
        file_path = _build_excel_to_server(rows, d)

        job2 = _read_job(job_id) or {}
        _write_job(job_id, {
            **job2,
            "status": "success",
            "message": "Експортот е успешен. Excel е снимен на сервер.",
            "file_path": file_path,
            "download_url": f"/siglife-report/Finansii/download/{job_id}",
            "finished_at": time.time()
        })
    except Exception as e:
        job2 = _read_job(job_id) or {}
        _write_job(job_id, {
            **job2,
            "status": "error",
            "message": f"Грешка при експорт: {str(e)}",
            "file_path": None,
            "download_url": None,
            "finished_at": time.time()
        })

@router.post("/Finansii/export-async")
async def finansii_export_async(
    request: Request,
    background_tasks: BackgroundTasks,
    export_date_form: str = Form(...),
    user=Depends(get_finansii_user)
):
    try:
        d = datetime.strptime(export_date_form, "%Y-%m-%d").date()
    except ValueError:
        return JSONResponse({"ok": False, "error": "Invalid date format. Use YYYY-MM-DD."}, status_code=400)

    job_id = str(uuid.uuid4())
    _write_job(job_id, {
        "ok": True,
        "type": "export",
        "status": "queued",
        "message": "",
        "file_path": None,
        "download_url": None,
        "started_at": time.time(),
        "finished_at": None
    })

    background_tasks.add_task(_export_job_runner_file, job_id, d)
    return JSONResponse({"ok": True, "job_id": job_id})

@router.get("/Finansii/export-job/{job_id}")
async def finansii_export_job_status(job_id: str, user=Depends(get_finansii_user)):
    job = _read_job(job_id)
    if not job:
        return JSONResponse({"ok": False, "error": "Job not found"}, status_code=404)
    return JSONResponse(job)

@router.get("/Finansii/download/{job_id}")
async def finansii_download(job_id: str, user=Depends(get_finansii_user)):
    job = _read_job(job_id)
    if not job or job.get("status") != "success":
        return Response("File not ready", status_code=404)

    file_path = job.get("file_path")
    if not file_path or not os.path.exists(file_path):
        return Response("File missing on server", status_code=404)

    return FileResponse(
        file_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=os.path.basename(file_path),
    )
