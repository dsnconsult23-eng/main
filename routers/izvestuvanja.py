from fastapi import APIRouter, Request, Form, Depends, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, FileResponse
from fastapi.templating import Jinja2Templates
from fastapi.responses import StreamingResponse

from datetime import datetime, date
from uuid import uuid4
import os
from typing import Optional, List, Tuple
import smtplib
from email.message import EmailMessage
from email.utils import make_msgid, formatdate
import time
import requests
import base64

import threading
import traceback
from pathlib import Path

# PDF (ReportLab)
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from routers import Connection
from .GeneratePDFOpomeni import generate_pdf
import pandas as pd
import io
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfWriter
from pydantic import BaseModel
from routers import SendMailLog 
import json

router = APIRouter()
templates = Jinja2Templates(directory=os.path.join(os.path.dirname(__file__), "..", "templates"))

CHANNEL_MAP = {
    "BANK": 33,
    "DIRECT": 57,
    "AGENCY_COMP": 58,
    "AGENTS": 77,
    "BROKER_COMP": 78,
    "PROMOTOR": 775
}

# ===========================
# Session check (HTML pages)
# ===========================
def get_current_user(request: Request):
    user = request.session.get("user")
    if not user:
        return RedirectResponse(url="/siglife-report/login", status_code=303)
    return user

# ===========================
# Session check (API routes -> JSON 401, NOT HTML redirect)
# ===========================
def get_current_user_api(request: Request):
    user = request.session.get("user")
    if not user:
        raise HTTPException(status_code=401, detail="Not authenticated")
    return user

# ===========================
# Pages
# ===========================

@router.get("/izvestuvanja", response_class=HTMLResponse)
async def izvestuvanja_page(request: Request, user=Depends(get_current_user)):
    return templates.TemplateResponse(
        "izvestuvanja.html",
        {
            "request": request,
            "user": user,
            "target_type": "physical",
            "subject": "Известување за промена на називот на друштвото за животно осигурување",
            "body": "",
            "attach_current_state": False,
            "message": None,

            # (опционално) ако во template сакаш да пополниш default за опомени
            "op_target_type": "all",
            "policy_no": "",
            "date_from": "",
            "date_to": "",
            "op_message": None,
        }
    )

CURRENT_STATE_PDF_PATH = "/opt/siglife-reporting/Image/Tekovna_sostojba_SigalLife.pdf"

# ===========================
# TEST MODE CONFIG
# ===========================
TEST_MODE = False               # ← само ова менуваш
TEST_EMAIL = ""
TEST_MAX_EMAILS = 5            # колку маилови да прати по тип

# ===========================
# DB cursor helper (Connection or Cursor)
# ===========================
def get_cursor_from_osis():
    obj, ok = Connection.OSISinit()
    if not ok:
        raise Exception("❌ Error - no connection to the database")
    print("OSISinit type:", type(obj))

    try:
        cur = obj.cursor()
        return cur, obj        # cursor, connection
    except Exception:
        return obj, None       # cursor, no connection

# ===========================
# Simple in-memory jobs store (shared for Izvestuvanja + Opomeni)
# ===========================
JOBS = {}  # job_id -> dict(...)
def _set_job(job_id: str, **kwargs):
    job = JOBS.get(job_id, {})
    job.update(kwargs)
    JOBS[job_id] = job

# ===========================
# NEW: request model for selected mails
# ===========================
class SendSelectedMailRequest(BaseModel):
    ids: list[str]
    job_id: Optional[str] = None
    test_email: Optional[str] = None

# ===========================
# SMTP settings (read once per job)
# ===========================
def get_smtp_settings():
    sql = "SELECT key, value FROM adm_appsettings WHERE key LIKE 'smtp.%'"

    cur, conn = get_cursor_from_osis()
    try:
        cur.execute(sql)
        settings = {}
        while True:
            row = cur.fetchone()
            if not row:
                break
            settings[str(row[0])] = row[1]
        return settings
    finally:
        try:
            cur.close()
        except Exception:
            pass
        if conn:
            try:
                conn.close()
            except Exception:
                pass

def load_smtp_config_once():
    settings = get_smtp_settings()

    sender_email = settings.get("smtp.fromaddress")
    sender_username = settings.get("smtp.username")
    sender_password = settings.get("smtp.password")

    if not sender_email or not sender_username or not sender_password:
        raise Exception("Missing SMTP credentials (smtp.fromaddress / smtp.username / smtp.password)")

    return {
        "sender_email": sender_email,
        "sender_username": sender_username,
        "sender_password": sender_password,
        "smtp_host": "smtp.office365.com",
        "smtp_port": 587
    }

# ===========================
# Recipients (SQL unchanged)
# ===========================
def fetch_recipients(target_type: str):
    sql_physical = """
    SELECT     
    DISTINCT LOWER(TRIM(email)) AS email
    FROM (
        SELECT c.email AS email
        FROM os_polisa a
        JOIN os_ponuda b ON a.os_ponudaid = b.os_ponudaid
        JOIN par_client c ON b.dogovoruvac_par_client = c.par_clientid
        WHERE b.par_statusid IN (17,18)
          AND a.polisa_pod_broj = vratipodbrojpolisa(a.polisa_broj, b.os_produktid)
          AND c.client_tip_pf = 'F'
          AND c.email IS NOT NULL
          AND TRIM(c.email) <> ''
          AND c.email LIKE '%@%.%'

        UNION

        SELECT c.email AS email
        FROM os_polisa a
        JOIN os_ponuda b ON a.os_ponudaid = b.os_ponudaid
        JOIN par_client c ON b.osigurenik_par_client = c.par_clientid
        WHERE b.par_statusid IN (17,18)
          AND a.polisa_pod_broj = vratipodbrojpolisa(a.polisa_broj, b.os_produktid)
          AND c.client_tip_pf = 'F'
          AND c.email IS NOT NULL
          AND TRIM(c.email) <> ''
          AND c.email LIKE '%@%.%'
    ) t
    where  LOWER(TRIM(email)) not in ( select recipient from email_audit_log)
    ORDER BY 1
    """

    sql_legal = """
    SELECT DISTINCT LOWER(TRIM(c.email)) AS email
    FROM os_polisa a
    JOIN os_ponuda b ON a.os_ponudaid = b.os_ponudaid
    JOIN par_client c ON b.dogovoruvac_par_client = c.par_clientid
    WHERE b.par_statusid IN (17,18)
      AND a.polisa_pod_broj = vratipodbrojpolisa(a.polisa_broj, b.os_produktid)
      AND c.client_tip_pf = 'P'
      AND c.email IS NOT NULL
      AND TRIM(c.email) <> ''
      AND c.email LIKE '%@%.%'
      AND  LOWER(TRIM(C.email)) not in ( select recipient from email_audit_log)
    ORDER BY 1
    """

    sql_all = """
    SELECT DISTINCT LOWER(TRIM(email)) AS email
    FROM (
        SELECT c.email AS email
        FROM os_polisa a
        JOIN os_ponuda b ON a.os_ponudaid = b.os_ponudaid
        JOIN par_client c ON b.dogovoruvac_par_client = c.par_clientid
        WHERE b.par_statusid IN (17,18)
          AND a.polisa_pod_broj = vratipodbrojpolisa(a.polisa_broj, b.os_produktid)
          AND c.email IS NOT NULL
          AND TRIM(c.email) <> ''
          AND c.email LIKE '%@%.%'

        UNION

        SELECT c.email AS email
        FROM os_polisa a
        JOIN os_ponuda b ON a.os_ponudaid = b.os_ponudaid
        JOIN par_client c ON b.osigurenik_par_client = c.par_clientid
        WHERE b.par_statusid IN (17,18)
          AND a.polisa_pod_broj = vratipodbrojpolisa(a.polisa_broj, b.os_produktid)
          AND c.email IS NOT NULL
          AND TRIM(c.email) <> ''
          AND c.email LIKE '%@%.%'
    ) t
    ORDER BY 1
    """

    sql = sql_physical if target_type == "physical" else sql_legal if target_type == "legal" else sql_all

    cur, conn = get_cursor_from_osis()
    try:
        cur.execute(sql)
        rows = cur.fetchall()
    finally:
        try:
            cur.close()
        except Exception:
            pass
        if conn:
            try:
                conn.close()
            except Exception:
                pass

    return [{"email": r[0]} for r in (rows or []) if r and r[0]]

# ===========================
# Audit log
# ===========================
def _audit_email_to_db(
    *,
    job_id: Optional[str],
    recipient: str,
    subject: str,
    message_id: Optional[str],
    status: str,
    error_message: Optional[str],
    attachments_cnt: int,
):
    results, OK = Connection.OSISinit()
    if not OK:
        print("[AUDIT][DB] ❌ No DB connection")
        return

    cur = results  # кај тебе најчесто ова е cursor
    try:
        sql = """
            INSERT INTO email_audit_log
                (job_id, recipient, subject, message_id, status, error_message, attachments_cnt, created_at)
            VALUES
                (?, ?, ?, ?, ?, ?, ?, CURRENT YEAR TO SECOND)
        """
        cur.execute(sql, (
            job_id,
            recipient,
            subject,
            message_id,
            status,
            error_message,
            int(attachments_cnt),
        ))

        try:
            cur.connection.commit()
        except Exception:
            pass

    except Exception as e:
        print("[AUDIT][DB] ❌ Insert failed:", repr(e))

# ===========================
# Send many emails using ONE SMTP connection
# ===========================
def send_many_emails_office365(
    emails: List[str],
    subject: str,
    html_body: str,
    attachments: List[Tuple[str, str, bytes]],
    smtp_cfg: dict,
    job_id: Optional[str] = None,
    max_retries_per_email: int = 3,
    reconnect_backoff_sec: float = 2.0,
):
    sent = 0
    failed = 0

    host = smtp_cfg.get("smtp_host")
    port = smtp_cfg.get("smtp_port")
    username = smtp_cfg.get("sender_username")
    password = smtp_cfg.get("sender_password")
    from_email = smtp_cfg.get("sender_email")

    attachments = attachments or []

    def connect_smtp() -> smtplib.SMTP:
        srv = smtplib.SMTP(host, port, timeout=60)
        srv.ehlo()
        srv.starttls()
        srv.ehlo()
        srv.login(username, password)
        return srv

    def build_msg(to_email: str) -> EmailMessage:
        msg = EmailMessage()
        msg["From"] = from_email
        msg["To"] = to_email
        msg["Subject"] = subject
        msg["Message-ID"] = make_msgid()
        msg["Date"] = formatdate(localtime=True)

        msg.set_content("This email contains HTML content.")
        msg.add_alternative(html_body or "", subtype="html")

        for filename, mime_type, content in attachments:
            if not filename or not content:
                continue
            if mime_type and "/" in mime_type:
                maintype, subtype = mime_type.split("/", 1)
            else:
                maintype, subtype = "application", "octet-stream"

            msg.add_attachment(
                content,
                maintype=maintype,
                subtype=subtype,
                filename=filename
            )
        return msg

    server = None
    try:
        server = connect_smtp()

        for idx, email in enumerate(emails, start=1):
            last_err = None

            for attempt in range(1, max_retries_per_email + 1):
                try:
                    if server is None:
                        server = connect_smtp()

                    msg = build_msg(email)
                    server.send_message(msg)
                    sent += 1

                    try:
                        _audit_email_to_db(
                            job_id=job_id,
                            recipient=email,
                            subject=subject,
                            message_id=str(msg["Message-ID"]),
                            status="SENT",
                            error_message=None,
                            attachments_cnt=len(attachments),
                        )
                    except Exception:
                        pass

                    last_err = None
                    break

                except (smtplib.SMTPServerDisconnected, smtplib.SMTPConnectError) as e:
                    last_err = repr(e)

                    try:
                        if server is not None:
                            server.quit()
                    except Exception:
                        pass
                    server = None

                    time.sleep(reconnect_backoff_sec * attempt)
                    continue

                except smtplib.SMTPResponseException as e:
                    last_err = f"{e.smtp_code} {e.smtp_error!r}"
                    if 400 <= e.smtp_code < 500:
                        time.sleep(reconnect_backoff_sec * attempt)
                        continue
                    else:
                        break

                except Exception as e:
                    last_err = repr(e)
                    break

            if last_err:
                failed += 1
                err_msg = f"{email}: {last_err}"

                try:
                    _audit_email_to_db(
                        job_id=job_id,
                        recipient=email,
                        subject=subject,
                        message_id=None,
                        status="FAILED",
                        error_message=last_err,
                        attachments_cnt=len(attachments),
                    )
                except Exception:
                    pass

                if job_id:
                    try:
                        _set_job(job_id, last_error=err_msg)
                    except Exception:
                        pass

    finally:
        try:
            if server is not None:
                server.quit()
        except Exception:
            pass

    return sent, failed

# ===========================
# Izvestuvanja: start send job
# ===========================
@router.post("/izvestuvanja/send-async")
async def send_async(
    request: Request,
    user=Depends(get_current_user_api),
    target_type: str = Form(...),  # physical | legal | all
    subject: str = Form(...),
    body: str = Form(...),
    attach_current_state: Optional[str] = Form(None),
    file: Optional[UploadFile] = File(None)
):
    job_id = str(uuid4())
    _set_job(job_id, status="queued", message="Queued", created_at=str(datetime.now()))

    attach_cs = True if attach_current_state is not None else False
    if target_type == "legal":
        attach_cs = True

    extra_attachment = None
    if file and file.filename:
        content = await file.read()
        extra_attachment = (file.filename, file.content_type or "application/octet-stream", content)

    try:
        _set_job(job_id, status="running", message="Sending...")

        recipients = fetch_recipients(target_type)
        emails = [(r.get("email") or "").strip() for r in recipients]
        emails = [e for e in emails if e]

        # dedup
        emails = list(dict.fromkeys(emails))

        if TEST_MODE:
            original_count = len(emails)
            if not emails:
                emails = [TEST_EMAIL]
            else:
                emails = emails[:TEST_MAX_EMAILS]
            print(f"[TEST MODE] target={target_type} | original={original_count} | sending={len(emails)}")

        attachments: List[Tuple[str, str, bytes]] = []

        if attach_cs:
            if os.path.exists(CURRENT_STATE_PDF_PATH) and os.path.getsize(CURRENT_STATE_PDF_PATH) > 0:
                with open(CURRENT_STATE_PDF_PATH, "rb") as f:
                    attachments.append(("Tekovna_sostojba_SigalLife.pdf", "application/pdf", f.read()))
            else:
                _set_job(job_id, last_error=f"❌ Missing Tekovna_sostojba PDF: {CURRENT_STATE_PDF_PATH}")

        if extra_attachment:
            attachments.append(extra_attachment)

        html_body = body
        if "<html" not in (body or "").lower() and "<br" not in (body or "").lower():
            html_body = "<pre style='font-family:Arial;white-space:pre-wrap;'>" + (body or "") + "</pre>"

        smtp_cfg = load_smtp_config_once()

        # ако сакаш секогаш фикс subject:
        subject1 = (
            "Информација: "
            "Известување за промена на називот на друштвото за животно осигурување "
        )

        sent, failed = send_many_emails_office365(
            emails=emails,
            subject=subject1,
            html_body=html_body,
            attachments=attachments,
            smtp_cfg=smtp_cfg,
            job_id=job_id
        )

        _set_job(
            job_id,
            status="success",
            message="Испраќањето е завршено.",
            sent_count=sent,
            failed_count=failed
        )

    except Exception as e:
        _set_job(job_id, status="failed", message=f"Грешка: {e}", last_error=repr(e))

    return JSONResponse({"ok": True, "job_id": job_id})

# ===========================
# Izvestuvanja: poll job status
# ===========================
@router.get("/izvestuvanja/job/{job_id}")
async def job_status(request: Request, job_id: str, user=Depends(get_current_user_api)):
    job = JOBS.get(job_id)
    if not job:
        return JSONResponse({"ok": False, "error": "Непознат job_id"})

    return JSONResponse({
        "ok": True,
        "status": job.get("status"),
        "message": job.get("message"),
        "sent_count": job.get("sent_count"),
        "failed_count": job.get("failed_count"),
        "last_error": job.get("last_error")
    })

# ======================================================================================
# =============================== OPOMENI (PDF) ========================================
# ======================================================================================

OPOMENI_OUT_DIR = Path("/opt/siglife-reporting/UNIQA/opomeni")
OPOMENI_OUT_DIR.mkdir(parents=True, exist_ok=True)

def build_opomeni_pdf(
    pdf_path: Path,
    *,
    target_type: str,
    date_from: date,
    date_to: date,
    policy_no: Optional[str],
    generated_by: str = ""
):
    c = canvas.Canvas(str(pdf_path), pagesize=A4)
    w, h = A4

    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, h - 60, "ОПОМЕНИ (PDF)")

    c.setFont("Helvetica", 11)
    c.drawString(50, h - 95, f"Тип клиенти: {target_type}")
    c.drawString(50, h - 115, f"Период: {date_from.isoformat()} до {date_to.isoformat()}")
    c.drawString(50, h - 135, f"Полиса број: {policy_no or '— (сите)'}")
    c.drawString(50, h - 155, f"Генерирано од: {generated_by or '—'}")
    c.drawString(50, h - 175, f"Датум: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    # TODO: тука ќе ставиш реални ставки
    y = h - 220
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "Пример ставки (замени со реални податоци):")
    y -= 20

    c.setFont("Helvetica", 10)
    sample_rows = [
        ("Клиент", "Полиса", "Износ", "Рок"),
        ("Име Презиме", policy_no or "S0197690", "1,250.00 MKD", "2026-01-31"),
        ("Фирма ДОО", "P-000999", "9,800.00 MKD", "2026-02-05"),
    ]
    for row in sample_rows:
        c.drawString(50, y, f"{row[0]:<18}  {row[1]:<12}  {row[2]:<14}  {row[3]}")
        y -= 16
        if y < 70:
            c.showPage()
            y = h - 60

    c.showPage()
    c.save()

def run_opomeni_job(
    job_id: str,
    *,
    target_type: str,
    df: date,
    dt: date,
    policy_no: Optional[str],
    client_id: Optional[str],
    sales_channels: Optional[List[str]],
    user_name: str
):
    try:
        _set_job(job_id, status="running", message="Се генерираат PDF опомени...")

        result = generate_pdfs_izvestuvanje_all(
            date_from=df.isoformat(),
            date_to=dt.isoformat(),
            target_type=target_type,
            sales_channels=sales_channels,
            policy_no=policy_no,
            client_id=client_id,
            save_folder=str(OPOMENI_OUT_DIR)
        )

        if not result.get("ok"):
            raise Exception(result.get("error", "Unknown error"))

        generated = result.get("generated", 0)
        processed = result.get("processed", 0)
        skipped = result.get("skipped", 0)
        errors = result.get("errors", 0)
        folder = result.get("folder")
        merged_pdf = result.get("merged_pdf")
        items = result.get("items", [])

        _set_job(
            job_id,
            status="ready",
            message=f"Генерирани {generated} PDF опомени.",
            generated=generated,
            processed=processed,
            skipped=skipped,
            errors=errors,
            folder=folder,
            file_path=merged_pdf,
            filename="Site_Opomeni.pdf",
            items=items
        )

    except Exception as e:
        _set_job(
            job_id,
            status="error",
            message="Грешка при генерирање PDF опомени.",
            last_error=str(e),
            trace=traceback.format_exc()
        )


@router.post("/opomeni/generate-async")
async def opomeni_generate_async(
    request: Request,
    user=Depends(get_current_user_api)
):
    try:
        form = await request.form()

        target_type = (form.get("target_type") or "all").strip()
        date_from = (form.get("date_from") or "").strip()
        date_to = (form.get("date_to") or "").strip()
        policy_no = (form.get("policy_no") or "").strip() or None
        client_id = (form.get("client_id") or "").strip() or None
        format_ = (form.get("format") or "pdf").strip()

        sales_channels = form.getlist("sales_channels")
        if not sales_channels:
            sales_channels = None

        if not date_from:
            return JSONResponse({"ok": False, "error": "date_from is required"})

        if not date_to:
            date_to = date_from

        try:
            df = datetime.strptime(date_from, "%Y-%m-%d").date()
            dt = datetime.strptime(date_to, "%Y-%m-%d").date()
        except Exception:
            return JSONResponse({"ok": False, "error": "Невалиден датум (очекувам YYYY-MM-DD)."})

        if dt < df:
            return JSONResponse({"ok": False, "error": "Датум 'до' не смее да е помал од 'од'."})

        job_id = str(uuid4())
        _set_job(
            job_id,
            status="queued",
            message="Queued",
            created_at=str(datetime.now()),
            file_path=None,
            filename=None,
            items=[]
        )

        user_name = ""
        try:
            if isinstance(user, dict):
                user_name = user.get("username") or user.get("name") or ""
            else:
                user_name = str(user)
        except Exception:
            user_name = ""

        t = threading.Thread(
            target=run_opomeni_job,
            kwargs=dict(
                job_id=job_id,
                target_type=target_type,
                df=df,
                dt=dt,
                policy_no=policy_no,
                client_id=client_id,
                sales_channels=sales_channels,
                user_name=user_name
            ),
            daemon=True
        )
        t.start()

        return JSONResponse({"ok": True, "job_id": job_id})

    except Exception as e:
        return JSONResponse({"ok": False, "error": str(e)}, status_code=500)



@router.get("/opomeni/download/{job_id}")
async def opomeni_download(request: Request, job_id: str, user=Depends(get_current_user_api)):
    job = JOBS.get(job_id)
    if not job:
        return JSONResponse({"ok": False, "error": "Непознат job_id"})

    if job.get("status") != "ready" or not job.get("file_path"):
        return JSONResponse({"ok": False, "error": "PDF не е подготвен."})

    pdf_path = Path(job["file_path"])
    if not pdf_path.exists():
        return JSONResponse({"ok": False, "error": "PDF фајлот не постои."})

    return FileResponse(
        str(pdf_path),
        media_type="application/pdf",
        filename=job.get("filename") or "Site_Opomeni.pdf"
    )

def norm_client_type(target_type: str) -> str:
    """
    UI -> DB mapping:
      physical -> F
      legal    -> P
      all      -> ALL
    """
    t = (target_type or "").lower().strip()
    if t in ("physical", "f", "fizicko", "fizicki"):
        return "F"
    if t in ("legal", "p", "pravno", "pravni"):
        return "P"
    return "ALL"


def generate_pdfs_izvestuvanje_all(
    *,
    date_from: str,
    date_to: str,
    target_type: str = "all",
    sales_channels: Optional[list[str]] = None,
    policy_no: Optional[str] = None,
    client_id: Optional[str] = None,
    save_folder: str = str(OPOMENI_OUT_DIR),
    limit: Optional[int] = None
) -> dict:

    print("========== START generate_pdfs_izvestuvanje_all ==========")

    processed = 0
    generated = 0
    skipped = 0
    errors = 0
    items = []

    # parse dates
    df = datetime.strptime(date_from, "%Y-%m-%d").date()
    dt = datetime.strptime(date_to, "%Y-%m-%d").date() if date_to else date.today()

    # Informix literal
    cutoff_date_sql = f"MDY({dt.month},{dt.day},{dt.year})"
    from_date_sql = f"MDY({df.month},{df.day},{df.year})"

    # client filter
    tt = norm_client_type(target_type)
    if tt == "F":
        client_filter = "AND pc.client_tip_pf = 'F'"
    elif tt == "P":
        client_filter = "AND pc.client_tip_pf = 'P'"
    else:
        client_filter = ""

    # channel filter
    channel_filter = ""
    if sales_channels:
        ids = [str(CHANNEL_MAP[c]) for c in sales_channels if c in CHANNEL_MAP]
        if ids:
            channel_filter = f"AND x3.par_prod_kanalid IN ({','.join(ids)})"

    # policy filter
    policy_filter = ""
    if policy_no and policy_no.strip():
        pol = esc_sql(policy_no.strip())
        policy_filter = f"AND x2.polisa_broj_cel = '{pol}'"

    # client_id filter
    clientid_filter = ""
    if client_id and str(client_id).strip():
        cid = esc_sql(str(client_id).strip())
        clientid_filter = f"AND pc.par_client = '{cid}'"

    today_str = datetime.today().strftime("%Y-%m-%d")
    full_save_path = os.path.join(save_folder, today_str)
    os.makedirs(full_save_path, exist_ok=True)

    merged_file_name = "Site_Opomeni.pdf"
    merged_full_path = os.path.join(full_save_path, merged_file_name)

    print("Saving PDFs to:", full_save_path)

    # ---------------------------------------------------------
    # 1) MAIN SQL - најди latest unpaid polisi / fakturi
    # ---------------------------------------------------------
    sql = f"""
    WITH base AS (
        SELECT
            x0.os_aneks_fakturaid,
            x2.polisa_broj_cel AS polisa_broj,
            x0.data_faktura,
            NVL(x0.iznos,0) AS iznos,

            pc.par_client        AS client_id,
            pc.desc              AS ime_prezime,
            pc.adresa_izvest_ulica AS adresa_grad,
            pc.email,
            pc.client_tip_pf     AS client_pf,

            UPPER((x1.os_aneks || '/' || TRIM(vesna.vrati_godina(x1.par_yearid)) || '-' || x0.rata)) AS faktura_broj,
            x1.par_clientid,
            x5.otp_code

        FROM viki.os_aneks_faktura x0
        JOIN viki.os_aneks   x1 ON x1.os_aneksid   = x0.os_aneksid
        JOIN viki.os_polisa  x2 ON x2.os_polisaid  = x1.os_polisaid
        JOIN viki.os_ponuda  x3 ON x3.os_ponudaid  = x2.os_ponudaid
        JOIN viki.os_produkt x4 ON x4.os_produktid = x3.os_produktid
        JOIN vesna.os_tipprodukt x5 ON x5.os_tipproduktid = x4.os_tipproduktid
        JOIN par_client pc ON pc.par_clientid = x1.par_clientid

        WHERE NVL(x0.f_rs,'R') <> 'N'
          AND x1.os_zbiren_aneksid IS NULL
          AND x3.skadenca_datum_do > TODAY
          
          AND {cutoff_date_sql} >= CASE
              WHEN x5.otp_code IN ('КО','UL','РИ','РК')
                   THEN DATE(ADD_MONTHS(x0.data_faktura, 2))
              ELSE DATE(ADD_MONTHS(x0.data_faktura, 5))
          END

          {client_filter}
          {channel_filter}
          {policy_filter}
          {clientid_filter}
    ),
    naplata AS (
        SELECT
            fs.os_aneks_fakturaid,
            SUM(NVL(fs.iznos_p,0)) AS naplata
        FROM fin_stavka fs
        JOIN base b ON b.os_aneks_fakturaid = fs.os_aneks_fakturaid
        WHERE fs.sifra_zatvaranje IS NOT NULL
          AND fs.f_rs <> 'N'
        GROUP BY fs.os_aneks_fakturaid
    ),
    fakturi AS (
        SELECT
            b.polisa_broj,
            b.client_id,
            b.ime_prezime,
            b.adresa_grad,
            b.email,
            b.faktura_broj,
            b.data_faktura,
            b.par_clientid,
            SUM(b.iznos - NVL(n.naplata,0)) AS saldo
        FROM base b
        LEFT JOIN naplata n ON n.os_aneks_fakturaid = b.os_aneks_fakturaid
        GROUP BY
            b.polisa_broj, b.client_id, b.ime_prezime, b.adresa_grad, b.email,
            b.faktura_broj, b.data_faktura, b.par_clientid
    )
       ,
     polisa_max AS (
    SELECT
        a.polisa_broj_cel,
        MAX(a.polisa_pod_broj) AS max_polisa_pod_broj
    FROM os_polisa a
    GROUP BY a.polisa_broj_cel
),
polisa_info AS (
    SELECT
        a.polisa_broj_cel,
        NVL(a.polisa_zamena, a.polisa_broj_cel) AS polisa_zamena,
        a.polisa_pod_broj,
        vrati_status(b.par_statusid) AS ponuda_status,
         a.status_polisa
    FROM os_polisa a
    JOIN polisa_max pm
        ON pm.polisa_broj_cel = a.polisa_broj_cel
       AND pm.max_polisa_pod_broj = a.polisa_pod_broj
    JOIN os_ponuda b
        ON a.os_ponudaid = b.os_ponudaid
)
         SELECT polisa_broj,par_clientid
         FROM fakturi f join  polisa_info pi on f.polisa_broj=pi.polisa_broj_cel
         where pi.status_polisa='K'
         GROUP BY polisa_broj,par_clientid
         HAVING SUM(saldo) > 0
    
    """

    print("Executing main SQL...")
    print(sql)

    try:
        results, OK = Connection.OSISinit()
        if not OK:
            return {"ok": False, "error": "Database connection failed"}

        results.execute(sql)

        selected_rows = []
        client_ids = set()
        policy_numbers = set()


        while True:
            row = results.fetchone()
            if not row:
                break

            processed += 1

            if limit and processed > limit:
                print("Limit reached:", limit)
                break

            polisa_broj, par_clientid = row
            print(f"Main row {processed}: par_clientid={par_clientid}, polisa={polisa_broj}")

            selected_rows.append({
                "par_clientid": par_clientid,
                "polisa_broj": polisa_broj
            })
            if par_clientid is not None:
                client_ids.add(int(par_clientid))
            
            if polisa_broj:
                policy_numbers.add(str(polisa_broj))

        if not selected_rows:
            print("⚠️ No rows from main SQL")
            return {
                "ok": True,
                "processed": processed,
                "generated": 0,
                "skipped": 0,
                "errors": 0,
                "folder": full_save_path,
                "merged_pdf": None,
                "message": "No data found.",
                "items": []
            }

        client_ids_sql = ",".join(str(x) for x in sorted(client_ids))
        policy_numbers_sql = ",".join("'" + x + "'" for x in sorted(policy_numbers))

        print(policy_numbers_sql)
        # policy_numbers_sql = ",".join(f"'{x.replace(\"'\", \"''\")}'" for x in sorted(policy_numbers))
        # ---------------------------------------------------------
        # 2) DETAIL SQL - врати ги деталите за PDF за тие клиенти
        # ---------------------------------------------------------
        sql_all_clients = f"""
        WITH
        b AS
        (
            SELECT
                a.par_clientid,
                a.par_yearid,
                f.rata,
                a.os_aneksid,
                a.br_rati,
                a.os_aneks,
                f.os_aneks_fakturaid,
                f.data_faktura AS datum_faktura,
                f.data_valuta  AS datum_valuta,
                f.iznos,
                f.iznos_denari AS iznosden,
                p.os_ponudaid,
                p.polisa_broj_cel,
                o.skadenca_datum_od,
                o.rok_plakanje,
                UPPER((a.os_aneks || '/' || TRIM(vesna.vrati_godina(a.par_yearid)) || '-' || f.rata)) AS faktura,
                ((a.os_aneks || a.par_yearid) || f.rata::INTEGER) AS report_id,
                TRIM(vesna.vrati_godina(a.par_yearid)) - YEAR(o.skadenca_datum_od) + 1 AS koja_godina,
                f.rata AS rata_broj
            FROM viki.os_aneks_faktura f
            JOIN viki.os_aneks a
                ON a.os_aneksid = f.os_aneksid
            JOIN viki.os_polisa p
                ON p.os_polisaid = a.os_polisaid
            JOIN viki.os_ponuda o
                ON o.os_ponudaid = p.os_ponudaid
            WHERE p.status_polisa = 'K'
              AND a.par_clientid IN ({client_ids_sql})
              AND p.polisa_broj_cel IN ({policy_numbers_sql})
              AND f.data_faktura < {cutoff_date_sql} 
        ),

        nbf2 AS
        (
            SELECT
                fs.os_aneks_fakturaid,
                NVL(SUM(NVL(fs.iznos_p, 0)), 0) AS naplata
            FROM viki.fin_stavka fs
            JOIN b
              ON fs.os_aneks_fakturaid = b.os_aneks_fakturaid
            WHERE fs.sifra_zatvaranje IS NOT NULL
              AND fs.f_rs != 'N'
            GROUP BY fs.os_aneks_fakturaid
        ),

        c AS
        (
            SELECT
                pc.par_clientid,
                pc.desc AS client_name,
                pc.client_tip_pf AS dogovoruvac_prav_fiz,
                pc.email
            FROM vesna.par_client pc
            WHERE pc.par_clientid IN ({client_ids_sql})
        ),

        s AS
        (
            SELECT
                b1.par_clientid,
                c.client_name,
                c.email,
                vesna.vrati_client_adresa_izv(b1.par_clientid) AS client_adresa,
                vesna.vrati_client_post_izv(b1.par_clientid)   AS client_post,
                b1.par_yearid,
                b1.rata,
                b1.os_aneksid,
                b1.br_rati,
                b1.os_aneks,
                b1.polisa_broj_cel,
                b1.polisa_broj_cel AS polisa_broj,
                b1.faktura,
                b1.datum_faktura,
                b1.datum_valuta,
                b1.skadenca_datum_od,
                b1.koja_godina,
                appuser.vrati_valuta_ponuda(b1.os_ponudaid) AS valuta,
                c.dogovoruvac_prav_fiz,
                vesna.vrati_godina(b1.par_yearid) AS godina,
                CASE
                    WHEN SUBSTR(b1.polisa_broj_cel, 1, 2) IN ('26', '28', '25', '27')
                        THEN ADD_MONTHS(DATE(b1.datum_faktura), 3)
                    ELSE ADD_MONTHS(DATE(b1.datum_faktura), 6)
                END AS trg_date,
                b1.iznos,
                NVL(n.naplata, 0) AS naplata
            FROM b b1
            JOIN c
                ON c.par_clientid = b1.par_clientid
            LEFT JOIN nbf2 n
                ON b1.os_aneks_fakturaid = n.os_aneks_fakturaid
        ),sx as 
         (
         SELECT
                 par_clientid,
                 client_name,
                 email,
                 client_adresa,
                 client_post,
                 par_yearid,
                 rata,
                 os_aneksid,
                 br_rati,
                 os_aneks,
                 polisa_broj_cel,
                 polisa_broj,
                 faktura,
                 datum_faktura,
                 datum_valuta,
                 skadenca_datum_od,
                 koja_godina,
                 valuta,
                 dogovoruvac_prav_fiz,
                 godina,
                 trg_date,
                 sum(s.iznos) iznos,
                 sum(NVL(s.naplata, 0)) AS naplata
             FROM s
             group by 1,2,3,4,5,6,7,8,9,10,11,12,13,14,15,16,17,18,19,20,21
         ),
         x AS
         (
             SELECT
                 s.*,
                 SUM(NVL(s.iznos,0)) OVER
                 (
                     PARTITION BY s.polisa_broj
                     ORDER BY s.datum_faktura, s.faktura
                     ROWS UNBOUNDED PRECEDING
                 ) AS kum_iznos,
                 SUM(NVL(s.naplata,0)) OVER
                 (
                     PARTITION BY s.polisa_broj
                     ORDER BY s.datum_faktura, s.faktura
                     ROWS UNBOUNDED PRECEDING
                 ) AS kum_naplata,
                 ROW_NUMBER() OVER
                 (
                     PARTITION BY s.polisa_broj
                     ORDER BY s.datum_faktura DESC, s.faktura DESC
                 ) AS rn_last
             FROM sx s
         ),
         polisa_max AS (
             SELECT
                 a.polisa_broj_cel,
                 MAX(a.polisa_pod_broj) AS max_polisa_pod_broj
             FROM os_polisa a
             GROUP BY a.polisa_broj_cel
         ),
         polisa_info AS (
             SELECT
                 a.polisa_broj_cel,
                 NVL(a.polisa_zamena, a.polisa_broj_cel) AS polisa_zamena,
                 a.polisa_pod_broj,
                 vrati_status(b.par_statusid) AS ponuda_status,
                 a.status_polisa
             FROM os_polisa a
             JOIN polisa_max pm
                 ON pm.polisa_broj_cel = a.polisa_broj_cel
             AND pm.max_polisa_pod_broj = a.polisa_pod_broj
             JOIN os_ponuda b
                 ON a.os_ponudaid = b.os_ponudaid
         )
         SELECT
             client_name,
             client_adresa,
             client_post AS client_grad,
             polisa_broj AS polisa_number,
             datum_valuta AS due_date,
             iznos AS premium_amount,
             naplata AS paid_premium,
             koja_godina AS godina,
             rata AS rata,
             ((TO_CHAR(x.datum_faktura, '%d/%m/%Y') || '-') || TO_CHAR(x.datum_valuta, '%d/%m/%Y')) AS period,
             ROUND((kum_iznos - kum_naplata) - (iznos - naplata), 2) AS unpaid_premium,
             valuta,
             ROUND((kum_iznos - kum_naplata), 2) AS vk_premija,
             par_clientid,
             faktura,
             email,
             (par_clientid || '_' || faktura) AS klient_faktura
         FROM x join  polisa_info pi on x.polisa_broj=pi.polisa_broj_cel
          where pi.status_polisa='K'
         and rn_last = 1
           AND (kum_iznos - kum_naplata) > 0
         ORDER BY client_name, polisa_broj
        """

        print("Executing detail SQL...")
        print(sql_all_clients)

        results_client, OK2 = Connection.OSISinit()
        if not OK2:
            return {
                "ok": False,
                "processed": processed,
                "generated": 0,
                "skipped": 0,
                "errors": 1,
                "folder": full_save_path,
                "message": "Client DB connection failed.",
                "items": []
            }

        results_client.execute(sql_all_clients)
        all_rows = results_client.fetchall()

        if not all_rows:
            print("⚠️ No detail rows")
            return {
                "ok": True,
                "processed": processed,
                "generated": 0,
                "skipped": 0,
                "errors": 0,
                "folder": full_save_path,
                "merged_pdf": None,
                "message": "No client detail data found.",
                "items": []
            }

        temp_pdf_files = []

        for client_row in all_rows:
            try:
                (
                    client_name,
                    client_adresa,
                    client_grad,
                    polisa_number,
                    due_date,
                    premium_amount,
                    paid_premium,
                    godina,
                    rata,
                    period,
                    unpaid_premium,
                    valuta,
                    vk_premija,
                    par_clientid,
                    faktura,
                    email,
                    klient_faktura
                ) = client_row

                print("Processing PDF:", klient_faktura)

                safe_faktura = str(faktura).replace("/", "-").replace("\\", "-")
                temp_file_name = f"Opomena_{par_clientid}_{safe_faktura}.pdf"
                temp_full_path = os.path.join(full_save_path, temp_file_name)

                premium_amount = premium_amount or 0
                paid_premium = paid_premium or 0
                unpaid_premium = unpaid_premium or 0
                vk_premija = vk_premija or 0

                balance = premium_amount - paid_premium

                generate_pdf(
                    temp_full_path,
                    client_name,
                    client_adresa,
                    client_grad,
                    polisa_number,
                    due_date,
                    premium_amount,
                    paid_premium,
                    balance,
                    godina,
                    rata,
                    period,
                    unpaid_premium,
                    vk_premija,
                    valuta
                )

                temp_pdf_files.append(temp_full_path)
                generated += 1
                print("✅ PDF OK:", temp_full_path)

                item_id = str(uuid4())
                items.append({
                    "id": item_id,
                    "client_id": str(par_clientid or ""),
                    "client_name": client_name or "",
                    "policy_no": polisa_number or "",
                    "email": email or "",
                    "pdf_path": temp_full_path,
                    "pdf_url": f"/siglife-report/opomeni/item-pdf/{item_id}",
                    "send_ready": bool(email and str(email).strip())
                })

            except Exception as e:
                print("❌ PDF error:", e)
                errors += 1

        if temp_pdf_files:
            print("Merging PDFs into one file...")
            writer = PdfWriter()

            for pdf_file in temp_pdf_files:
                writer.append(pdf_file)

            with open(merged_full_path, "wb") as fout:
                writer.write(fout)

            writer.close()
            print("✅ Merged PDF created:", merged_full_path)

            # НЕ ги бриши temp фајловите ако сакаш preview + selective send
            # for pdf_file in temp_pdf_files:
            #     try:
            #         os.remove(pdf_file)
            #     except Exception as e:
            #         print("⚠️ Could not remove temp file:", pdf_file, e)
        else:
            print("⚠️ No PDFs generated to merge.")

    except Exception as e:
        print("❌ MAIN ERROR:", e)
        errors += 1

    print("========== END generate_pdfs_izvestuvanje_all ==========")

    return {
        "ok": errors == 0,
        "processed": processed,
        "generated": generated,
        "skipped": skipped,
        "errors": errors,
        "folder": full_save_path,
        "merged_pdf": merged_full_path if os.path.exists(merged_full_path) else None,
        "items": items
    }



def esc_sql(s: str) -> str:
    return s.replace("'", "''")


@router.post("/opomeni/export-excel")
async def export_opomeni_excel(
    target_type: str = Form("all"),
    sales_channels: list[str] = Form(None),
    policy_no: str = Form(None),
    client_id: str = Form(None),
    date_to: str = Form(None)
):

    if not date_to:
        date_to = date.today().isoformat()

    year_to, month_to, day_to = date_to.split("-")

    client_filter = ""
    if target_type == "physical":
        client_filter = "AND pc.client_tip_pf='F'"
    elif target_type == "legal":
        client_filter = "AND pc.client_tip_pf='P'"

    channel_filter = ""
    if sales_channels:
        ids = [str(CHANNEL_MAP[c]) for c in sales_channels if c in CHANNEL_MAP]
        if ids:
            channel_filter = f"AND x3.par_prod_kanalid IN ({','.join(ids)})"

    if policy_no and policy_no.strip():
        pol = esc_sql(policy_no.strip())
        policy_filter = f"AND x2.polisa_broj_cel = '{pol}'"
    else:
        policy_filter = ""

    if client_id and str(client_id).strip():
        cid = esc_sql(str(client_id).strip())
        clientid_filter = f"AND pc.par_client = '{cid}'"
    else:
        clientid_filter = ""
    

    if not date_to:
        date_to = date.today().isoformat()

    y, m, d = date_to.split("-")
    cutoff_date_sql = f"MDY({int(m)},{int(d)},{int(y)})"   # пример MDY(2,28,2026)
        

    sql = f"""
    WITH base AS (
        SELECT
            x0.os_aneks_fakturaid,
            x2.polisa_broj_cel AS polisa_broj,
            x0.data_faktura,
            NVL(x0.iznos,0) AS iznos,

            pc.par_client AS client_id,
            pc.desc AS ime_prezime,
            pc.adresa_izvest_ulica AS adresa_grad,
            pc.email,
            pc.client_tip_pf,
            x5.otp_code

        FROM viki.os_aneks_faktura x0
        JOIN viki.os_aneks   x1 ON x1.os_aneksid   = x0.os_aneksid
        JOIN viki.os_polisa  x2 ON x2.os_polisaid  = x1.os_polisaid
        JOIN viki.os_ponuda  x3 ON x3.os_ponudaid  = x2.os_ponudaid
        JOIN viki.os_produkt x4 ON x4.os_produktid = x3.os_produktid
        JOIN vesna.os_tipprodukt x5 ON x5.os_tipproduktid = x4.os_tipproduktid
        JOIN par_client pc ON pc.par_clientid = x1.par_clientid

        WHERE NVL(x0.f_rs,'R') <> 'N'
          AND x1.os_zbiren_aneksid IS NULL
          AND x3.skadenca_datum_do > TODAY

          AND {cutoff_date_sql}  >= CASE
                WHEN x5.otp_code IN ('КО','UL','РИ','РК')
                     THEN DATE(ADD_MONTHS(x0.data_faktura,2))
                ELSE DATE(ADD_MONTHS(x0.data_faktura,5))
          END

          {client_filter}
          {channel_filter}
          {policy_filter}
          {clientid_filter}
    ),
    naplata AS (
        SELECT
            fs.os_aneks_fakturaid,
            SUM(NVL(fs.iznos_p,0)) AS naplata
        FROM fin_stavka fs
        JOIN base b
          ON b.os_aneks_fakturaid = fs.os_aneks_fakturaid
        WHERE fs.sifra_zatvaranje IS NOT NULL
          AND fs.f_rs <> 'N'
        GROUP BY fs.os_aneks_fakturaid
    )

    SELECT
        b.polisa_broj,
        b.client_id,
        b.ime_prezime,
        b.adresa_grad,
        b.email,
        SUM(b.iznos - NVL(n.naplata,0)) AS saldo

    FROM base b
    LEFT JOIN naplata n
      ON n.os_aneks_fakturaid = b.os_aneks_fakturaid

    GROUP BY
        b.polisa_broj,
        b.client_id,
        b.ime_prezime,
        b.adresa_grad,
        b.email

    HAVING SUM(b.iznos - NVL(n.naplata,0)) > 0
    """



    results, OK = Connection.OSISinit()

    if not OK:
        return {"ok": False}
    print(sql)
    print("SQL:", sql)

    results.execute(sql)

    rows = results.fetchall()
    cols = [c[0] for c in results.description]

    df = pd.DataFrame(rows, columns=cols)

    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:

        df.to_excel(writer, index=False, sheet_name="Opomeni")

        wb = writer.book
        ws = writer.sheets["Opomeni"]

        # bold header
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        # currency format
        saldo_col = None
        for i, name in enumerate(df.columns, start=1):
            if name.lower() == "saldo":
                saldo_col = i

        if saldo_col:
            for r in range(2, ws.max_row + 1):
                ws.cell(r, saldo_col).number_format = '#,##0.00'

        # auto column width
        for col_idx, col in enumerate(df.columns, 1):

            length = max(
                df[col].astype(str).map(len).max(),
                len(col)
            )

            ws.column_dimensions[get_column_letter(col_idx)].width = length + 2

        ws.freeze_panes = "A2"

    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=opomeni.xlsx"}
    )

@router.get("/opomeni/status/{job_id}")
async def opomeni_status(job_id: str, user=Depends(get_current_user_api)):
    job = JOBS.get(job_id)
    if not job:
        return JSONResponse({"ok": False, "error": "Непознат job_id"})

    file_url = None
    if job.get("status") == "ready":
        file_url = f"/siglife-report/opomeni/download/{job_id}"

    return JSONResponse({
        "ok": True,
        "status": job.get("status"),
        "message": job.get("message"),
        "generated": job.get("generated", 0),
        "processed": job.get("processed", 0),
        "skipped": job.get("skipped", 0),
        "errors": job.get("errors", 0),
        "file_url": file_url,
        "last_error": job.get("last_error"),
        "items": job.get("items", [])
    })

# ===========================
# NEW: preview individual generated pdf
# ===========================
@router.get("/opomeni/item-pdf/{item_id}")
async def opomeni_item_pdf(item_id: str, user=Depends(get_current_user_api)):
    for job in JOBS.values():
        items = job.get("items", [])
        for item in items:
            if item.get("id") == item_id:
                pdf_path = item.get("pdf_path")
                if not pdf_path or not os.path.exists(pdf_path):
                    return JSONResponse({"ok": False, "error": "PDF не постои."}, status_code=404)

                return FileResponse(
                    pdf_path,
                    media_type="application/pdf",
                    filename=os.path.basename(pdf_path)
                )

    return JSONResponse({"ok": False, "error": "Непознат item_id"}, status_code=404)

# ===========================
# NEW: send selected generated pdfs by mail
# ===========================
@router.post("/opomeni/send-selected-mail")
async def send_selected_opomeni_mail(
    payload: SendSelectedMailRequest,
    user=Depends(get_current_user_api)
):
    try:
        selected_ids = payload.ids or []
        if not selected_ids:
            return JSONResponse(
                {"ok": False, "error": "Нема селектирани записи."},
                status_code=400
            )

        smtp_cfg = load_smtp_config_once()

        sent_count = 0
        failed_count = 0

        test_email = (payload.test_email or "").strip()

        all_items = []
        for job in JOBS.values():
            all_items.extend(job.get("items", []))

        selected_items = [x for x in all_items if x.get("id") in selected_ids]

        if not selected_items:
            return JSONResponse(
                {"ok": False, "error": "Не се пронајдени селектираните записи."},
                status_code=404
            )

        for item in selected_items:
            try:
                # real_email = (item.get("email") or "").strip()
                real_email='svconsult40@gmail.com'
                email = test_email if test_email else real_email

                pdf_path = item.get("pdf_path")
                client_name = item.get("client_name") or ""
                policy_no = item.get("policy_no") or ""
                client_id = item.get("client_id")
                address = item.get("address")
                city = item.get("city")

                print("test_email =", test_email)
                print("real_email =", real_email)
                print("email used =", email)

                if not email:
                    failed_count += 1
                    continue

                if not pdf_path or not os.path.exists(pdf_path):
                    failed_count += 1
                    continue

                subject = f"Опомена за полиса {policy_no}"
                polisa_number = str(policy_no or "")

                loan_section = ""
                if not polisa_number.startswith("28"):
                    loan_section = """
                    <p>Подигнување на <a href="https://www.sigal.com.mk/подигнување-на-заем-по-полиса-за-осигу/"
                            style="color:#007bff; text-decoration:none;"><strong>ЗАЕМ</strong></a>
                        врз основ на Вашата полиса за осигурување на живот во рок од еден работен ден.</p>

                    <p>Сите информации околу условите и потребните документи за подигнување на заем,
                        може да ги најдете на следниот линк:
                        <a href="https://www.sigal.com.mk/подигнување-на-заем-по-полиса-за-осигу/" style="color:#007bff;">
                        Заем по полиса за осигурување на живот
                        </a>.
                    </p>
                    """

                html_body = f"""
                <html>
                    <body style="font-family: Arial, sans-serif; line-height: 1.6; color:#000;">

                    <p>Почитувани,</p>

                    <table cellpadding="0" cellspacing="0" border="0" width="100%" style="margin:20px 0;">
                    <tr>
                        <td width="120" valign="top" style="padding-right:15px;">
                        <img src="cid:uniqa_logo" alt="UNIQA" width="100" style="display:block;">
                        </td>

                        <td valign="top" style="font-family:Arial,sans-serif; font-size:14px; color:#000;">
                        <p style="margin:0 0 10px 0;">Во прилог Ви испраќаме опомена за неплатена премија за осигурување на живот.</p>

                        <p style="margin:0 0 10px 0;">Сега можете брзо и лесно да ја платите Вашата премија
                            <a href="https://plationline.sigal.com.mk/" style="color:#007bff; text-decoration:none;">онлајн</a>, без надоместок.
                        </p>

                        <p style="margin:0;"><strong>Ново!</strong> Премијата може да ја платите на рати преку кредитна картичка на
                            Стопанска Банка АД Скопје со кликање на копчето
                        </p>
                        </td>
                    </tr>
                    </table>

                    <p><strong>На рати преку „Стопанска Банка“ А.Д. Скопје</strong></p>

                    <p>СИГАЛ ЛАЈФ АД Скопје ви овозможува креирање на
                    <strong>траен налог</strong> за плаќање на премија за осигурување на живот со кликање на копчето</p>

                    <p>
                    <a href="https://plationline.sigal.com.mk/"
                        style="display:inline-block; padding:12px 25px; background:#007bff; color:#fff;
                                text-decoration:none; border-radius:6px; font-weight:bold;">
                        Плати премија онлајн
                    </a>
                    </p>

                    <p>Потврда за извршеното плаќање ќе добиете на Вашиот е-мејл.</p>

                    {loan_section}

                    <p>За дополнителни прашања или информации, слободно обратете се на телефонскиот број:
                    02/ 3 288 -820, E-mail:
                    <a href="mailto:lifeinsurance@sigal.com.mk" style="color:#007bff;">lifeinsurance@sigal.com.mk</a>,
                    Веб-страна: <a href="https://sigal.com.mk" style="color:#007bff;">www.sigal.com.mk</a>
                    </p>

                    <p style="font-size:12px; color:gray;">
                    Напомена: Оваа е-пошта и сите прилози содржат доверливи информации наменети само за употреба на лицето
                    на кое таа се однесува... Ако сте ја примиле оваа е-пошта по грешка, Ве молиме веднаш да нè известите
                    и потоа да ја избришете од Вашиот систем.
                    </p>

                    <p>Со почит,<br><strong>СИГАЛ ЛАЈФ АД Скопје</strong></p>
                    <hr>

                    <p style="font-size:12px; color:#333; line-height:1.4;">
                    SIGAL Life a.d. Skopje<br>
                    Bul. Ilinden br.1<br>
                    Skopje, Macedonia<br>
                    Tel.+389(0)23288820<br>
                    Fax.+389(0)23215128<br>
                    <a href="mailto:lifeinsurance@sigal.com.mk" style="color:#007bff;">lifeinsurance@sigal.com.mk</a><br>
                    <a href="https://sigal.com.mk" style="color:#007bff;">www.sigal.com.mk</a>
                    </p>

                </body>
                </html>
                """

                recipients = [
                    {
                        "email": email,
                        "client_id": client_id,
                        "policy_no": policy_no,
                        "client_name": client_name,
                        "address": address,
                        "city": city,
                        "pdf_paths": [pdf_path]
                    }
                ]

                sent, failed = SendMailLog.send_many_emails_with_logo_and_db_log(
                    recipients=recipients,
                    subject=subject,
                    html_body=html_body,
                    smtp_cfg=smtp_cfg,
                    job_id=None
                )

                sent_count += sent
                failed_count += failed

            except Exception as e:
                print("send_selected_opomeni_mail error:", e)
                failed_count += 1

        return JSONResponse({
            "ok": True,
            "message": "Испраќањето е завршено.",
            "sent_count": sent_count,
            "failed_count": failed_count
        })

    except Exception as e:
        return JSONResponse({"ok": False, "error": str(e)}, status_code=500)

@router.post("/opomeni/export-sent-emails-excel")
async def export_sent_emails_excel(
    request: Request,
    user=Depends(get_current_user_api)
):
    try:
        form = await request.form()

        date_from = (form.get("date_from") or "").strip()
        date_to = (form.get("date_to") or "").strip()

        print("date_from =", repr(date_from))
        print("date_to   =", repr(date_to))

        if not date_from:
            return JSONResponse({"ok": False, "error": "date_from is required"}, status_code=400)

        if not date_to:
            date_to = date_from

        try:
            df = datetime.strptime(date_from, "%Y-%m-%d").date()
            dt = datetime.strptime(date_to, "%Y-%m-%d").date()
        except ValueError as e:
            print("DATE PARSE ERROR:", e)
            return JSONResponse(
                {"ok": False, "error": f"Невалиден датум. date_from={date_from}, date_to={date_to}"},
                status_code=400
            )

        if dt < df:
            return JSONResponse({"ok": False, "error": "Датум 'до' не смее да е помал од 'од'."}, status_code=400)

        from datetime import timedelta
        dt_next = dt + timedelta(days=1)

        sql = f"""
            SELECT
                policy_no,
                client_id,
                client_name,
                address,
                city,
                recipient AS email
            FROM email_audit_log
            WHERE status = 'SENT' and tip='O'
              AND created_at >= MDY({df.month},{df.day},{df.year})
              AND created_at <  MDY({dt_next.month},{dt_next.day},{dt_next.year})
            ORDER BY created_at DESC
        """

        print(sql)

        cur, conn = get_cursor_from_osis()
        try:
            cur.execute(sql)
            rows = cur.fetchall()
            cols = [c[0] for c in cur.description]
        finally:
            try:
                cur.close()
            except Exception:
                pass
            if conn:
                try:
                    conn.close()
                except Exception:
                    pass

        df_excel = pd.DataFrame(rows, columns=cols if rows else [
            "policy_no", "client_id", "client_name", "address", "city", "email"
        ])

        df_excel = df_excel.rename(columns={
            "policy_no": "Број на полиса",
            "client_id": "ID на клиент",
            "client_name": "Име и презиме",
            "address": "Адреса на известување",
            "city": "Град",
            "email": "Меил"
        })

        output = io.BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df_excel.to_excel(writer, index=False, sheet_name="ИспратениОпомени")
            ws = writer.sheets["ИспратениОпомени"]

            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

            for col_idx, col in enumerate(df_excel.columns, 1):
                max_len = max(
                    len(str(col)),
                    max((len(str(v)) for v in df_excel[col].fillna("").tolist()), default=0)
                )
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)

            ws.freeze_panes = "A2"

        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=isprateni_opomeni.xlsx"}
        )

    except Exception as e:
        print("export_sent_emails_excel ERROR:", e)
        return JSONResponse({"ok": False, "error": str(e)}, status_code=500)

@router.post("/opomeni/send-kibs-test")
async def send_opomeni_kibs_test(
    payload: SendSelectedMailRequest,
    user=Depends(get_current_user_api)
):
    try:
        selected_ids = payload.ids or []

        if not selected_ids:
            return JSONResponse(
                {"ok": False, "error": "Нема селектирани записи."},
                status_code=400
            )

        sent_count = 0
        failed_count = 0
        job_id = str(uuid4())

        all_items = []
        for job in JOBS.values():
            all_items.extend(job.get("items", []))

        selected_items = [x for x in all_items if x.get("id") in selected_ids]

        if not selected_items:
            return JSONResponse(
                {"ok": False, "error": "Не се пронајдени записи."},
                status_code=404
            )

        for item in selected_items:
            try:
                pdf_path = item.get("pdf_path")
                email = item.get("email")
                client_name = item.get("client_name")
                client_id = item.get("client_id")
                policy_no = item.get("policy_no")
                address = item.get("address")
                city = item.get("city")

                if not pdf_path or not os.path.exists(pdf_path):
                    failed_count += 1

                    _audit_kibs_to_db(
                        job_id=job_id,
                        recipient=email or "",
                        subject=f"Опомена за полиса {policy_no or ''}",
                        status="FAILED",
                        error_message="PDF не постои",
                        client_id=client_id,
                        policy_no=policy_no,
                        client_name=client_name,
                        address=address,
                        city=city,
                        attachment_path=pdf_path,
                        attachment_filename=os.path.basename(pdf_path) if pdf_path else None,
                        attachment_position=1,
                        provider="KIBS_TEST"
                    )
                    continue

                with open(pdf_path, "rb") as f:
                    pdf_bytes = f.read()

                subject = f"Опомена за полиса {policy_no}"
                url = "https://splus-testapi.kibs.com.mk/v4/workflow/create"

                # според документацијата SignPlus REST API користи Basic HTTP auth
                # и workflow create е на /workflow/create под v4 base URL. :contentReference[oaicite:0]{index=0}

                auth = ("YOUR_KIBS_USERNAME", "YOUR_KIBS_PASSWORD")

                data_json = {
                    "name": subject,
                    "messageToAllSigners": "Опомена",
                    "signers": [
                        {
                            "signingProvider": "ON_THE_FLY",
                            "email": email,
                            "firstName": client_name or "Client",
                            "lastName": "."
                        }
                    ]
                }

                files = {
                    "file": (os.path.basename(pdf_path), pdf_bytes, "application/pdf")
                }

                response = requests.post(
                    url,
                    auth=auth,
                    files=files,
                    data={"data": json.dumps(data_json, ensure_ascii=False)},
                    timeout=60
                )

                response_text = response.text[:1000] if response.text else None
                external_id = None

                try:
                    response_json = response.json()
                    external_id = response_json.get("workflowUID")
                except Exception:
                    response_json = None

                if response.status_code == 200:
                    sent_count += 1

                    _audit_kibs_to_db(
                        job_id=job_id,
                        recipient=email or "",
                        subject=subject,
                        status="SENT",
                        error_message=None,
                        client_id=client_id,
                        policy_no=policy_no,
                        client_name=client_name,
                        address=address,
                        city=city,
                        attachment_path=pdf_path,
                        attachment_filename=os.path.basename(pdf_path),
                        attachment_position=1,
                        provider="KIBS_TEST",
                        external_id=external_id,
                        response_text=response_text
                    )

                else:
                    failed_count += 1

                    _audit_kibs_to_db(
                        job_id=job_id,
                        recipient=email or "",
                        subject=subject,
                        status="FAILED",
                        error_message=f"HTTP {response.status_code}",
                        client_id=client_id,
                        policy_no=policy_no,
                        client_name=client_name,
                        address=address,
                        city=city,
                        attachment_path=pdf_path,
                        attachment_filename=os.path.basename(pdf_path),
                        attachment_position=1,
                        provider="KIBS_TEST",
                        external_id=external_id,
                        response_text=response_text
                    )

            except Exception as e:
                failed_count += 1

                try:
                    _audit_kibs_to_db(
                        job_id=job_id,
                        recipient=item.get("email") or "",
                        subject=f"Опомена за полиса {item.get('policy_no') or ''}",
                        status="FAILED",
                        error_message=str(e),
                        client_id=item.get("client_id"),
                        policy_no=item.get("policy_no"),
                        client_name=item.get("client_name"),
                        address=item.get("address"),
                        city=item.get("city"),
                        attachment_path=item.get("pdf_path"),
                        attachment_filename=os.path.basename(item.get("pdf_path")) if item.get("pdf_path") else None,
                        attachment_position=1,
                        provider="KIBS_TEST"
                    )
                except Exception:
                    pass

        return JSONResponse({
            "ok": True,
            "message": "Испраќањето преку KIBS TEST е завршено.",
            "sent_count": sent_count,
            "failed_count": failed_count
        })

    except Exception as e:
        return JSONResponse(
            {"ok": False, "error": str(e)},
            status_code=500
        )
    
def _audit_kibs_to_db(
    *,
    job_id: Optional[str],
    recipient: str,
    client_id: Optional[str],
    policy_no: Optional[str],
    client_name: Optional[str],
    address: Optional[str],
    city: Optional[str],
    attachment_path: Optional[str],
    attachment_filename: Optional[str],
    kibs_transaction_id: Optional[str],
    status: str,
    error_message: Optional[str] = None
):

    results, OK = Connection.OSISinit()
    if not OK:
        print("[AUDIT][KIBS] ❌ No DB connection")
        return

    cur = results

    try:

        sql = """
        INSERT INTO kibs_audit_log
        (
            job_id,
            recipient,
            client_id,
            policy_no,
            client_name,
            address,
            city,
            attachment_path,
            attachment_filename,
            kibs_transaction_id,
            status,
            error_message,
            created_at
        )
        VALUES
        (
            ?,?,?,?,?,?,?,?,?,?,?,?,
            CURRENT YEAR TO SECOND
        )
        """

        cur.execute(sql, (
            job_id,
            recipient,
            client_id,
            policy_no,
            client_name,
            address,
            city,
            attachment_path,
            attachment_filename,
            kibs_transaction_id,
            status,
            error_message
        ))

        try:
            cur.connection.commit()
        except Exception:
            pass

    except Exception as e:
        print("[AUDIT][KIBS] ❌ Insert failed:", repr(e))