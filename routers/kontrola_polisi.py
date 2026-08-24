from fastapi import APIRouter, Request, Form, HTTPException
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from auth.role_utils import has_any_role
from routers import Connection
import os, io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

router = APIRouter()
templates = Jinja2Templates(
    directory=os.path.join(os.path.dirname(os.path.dirname(__file__)), "templates")
)

_ALLOWED_ROLES = ["admin", "kontrola_polisi"]


def _require_role(request: Request):
    user = request.session.get("user")
    if not user:
        return None, RedirectResponse(url="/siglife-report/login", status_code=303)
    if not has_any_role(user, *_ALLOWED_ROLES):
        raise HTTPException(status_code=403, detail="Немате пристап.")
    return user, None


def _run(sql, params=None):
    cursor, ok = Connection.OSISinit()
    if not ok or cursor is None:
        raise RuntimeError("Нема конекција кон базата")
    try:
        cursor.execute(sql, params or [])
        return cursor.fetchall()
    finally:
        try:
            cursor.close()
        except Exception:
            pass


# ------------------------------------------------------------------
# Проверка 1: полиси со статус 18 (капитализирана), 19, 20 (откуп), 36,
# 37 или 42 (раскината/сторнирана) — статусот се зема од ПОСЛЕДНАТА
# понуда (maxpodbroj_datum), со позитивно салдо (долг > 0).
#
# Салдото (износ - наплата) се пресметува по ИСТАТА логика како кај SMS
# (prebSMSDospeanaPremija во SMS.py): наплатата доаѓа од рачен join на
# fin_stavka (sifra_zatvaranje IS NOT NULL, f_rs<>'N'), НЕ преку
# vrati_naplata() UDF — намерно, за салдото тука да е конзистентно со
# она што го гледа/испраќа SMS модулот.
# ------------------------------------------------------------------
def check_kapitalizirana_faktura_nula():
    sql = """
        WITH included_polisi AS (
            SELECT DISTINCT p2.polisa_broj_cel
            FROM os_ponuda o2
            JOIN os_polisa p2 ON p2.os_ponudaid = o2.os_ponudaid
            WHERE o2.par_statusid IN (36, 37, 42, 18, 20, 19)
              AND o2.ponuda_podbroj = maxpodbroj_datum(o2.ponuda_broj, p2.polisa_broj, o2.os_produktid, TODAY)
        ),
        base AS (
            SELECT
                x0.os_aneks_fakturaid,
                TRIM(x2.polisa_broj_cel) AS polisa_broj_cel,
                NVL(x0.iznos, 0) AS iznos
            FROM os_aneks_faktura x0
            JOIN os_aneks  x1 ON x1.os_aneksid  = x0.os_aneksid
            JOIN os_polisa x2 ON x2.os_polisaid = x1.os_polisaid
            WHERE NVL(x0.f_rs, 'R') <> 'N'
              AND x1.os_zbiren_aneksid IS NULL
              AND x0.data_faktura <= TODAY
              AND EXISTS (
                  SELECT 1 FROM fin_stavka fg
                  WHERE fg.os_aneks_fakturaid = x0.os_aneks_fakturaid
                    AND fg.iznos_d IS NOT NULL
              )
        ),
        naplata AS (
            SELECT fs.os_aneks_fakturaid, SUM(NVL(fs.iznos_p, 0)) AS naplata
            FROM fin_stavka fs
            JOIN base b ON b.os_aneks_fakturaid = fs.os_aneks_fakturaid
            WHERE fs.sifra_zatvaranje IS NOT NULL
              AND fs.f_rs <> 'N'
            GROUP BY fs.os_aneks_fakturaid
        )
        SELECT
            b.polisa_broj_cel,
            MAX((
                SELECT CASE WHEN o3.par_polisa_statusid IS NULL THEN appuser.vrati_status(o3.par_statusid)
                            ELSE appuser.vrati_polisa_status_desc(o3.par_polisa_statusid) END
                FROM os_ponuda o3
                JOIN os_polisa p3 ON p3.os_ponudaid = o3.os_ponudaid
                WHERE p3.polisa_broj_cel = b.polisa_broj_cel
                  AND o3.ponuda_podbroj = maxpodbroj_datum(o3.ponuda_broj, p3.polisa_broj, o3.os_produktid, TODAY)
                  AND p3.os_polisaid = (
                      SELECT MAX(p3b.os_polisaid)
                      FROM os_ponuda o3b
                      JOIN os_polisa p3b ON p3b.os_ponudaid = o3b.os_ponudaid
                      WHERE p3b.polisa_broj_cel = b.polisa_broj_cel
                        AND o3b.ponuda_podbroj = maxpodbroj_datum(o3b.ponuda_broj, p3b.polisa_broj, o3b.os_produktid, TODAY)
                  )
            )) AS polisa_status_desc,
            MAX((
                SELECT appuser.vrati_status(o3.par_statusid)
                FROM os_ponuda o3
                JOIN os_polisa p3 ON p3.os_ponudaid = o3.os_ponudaid
                WHERE p3.polisa_broj_cel = b.polisa_broj_cel
                  AND o3.ponuda_podbroj = maxpodbroj_datum(o3.ponuda_broj, p3.polisa_broj, o3.os_produktid, TODAY)
                  AND p3.os_polisaid = (
                      SELECT MAX(p3b.os_polisaid)
                      FROM os_ponuda o3b
                      JOIN os_polisa p3b ON p3b.os_ponudaid = o3b.os_ponudaid
                      WHERE p3b.polisa_broj_cel = b.polisa_broj_cel
                        AND o3b.ponuda_podbroj = maxpodbroj_datum(o3b.ponuda_broj, p3b.polisa_broj, o3b.os_produktid, TODAY)
                  )
            )) AS polisa_sosotojba,
            SUM(b.iznos) AS vk_iznos,
            SUM(NVL(n.naplata, 0)) AS vk_naplata,
            SUM(b.iznos - NVL(n.naplata, 0)) AS saldo
        FROM base b
        LEFT JOIN naplata n ON n.os_aneks_fakturaid = b.os_aneks_fakturaid
        WHERE b.polisa_broj_cel IN (SELECT polisa_broj_cel FROM included_polisi)
        GROUP BY b.polisa_broj_cel
        HAVING ROUND(SUM(b.iznos - NVL(n.naplata, 0)), 2) > 0
        ORDER BY 1
    """
    rows = _run(sql)
    return [
        {
            "polisa_broj": r[0],
            "polisa_status_desc": r[1] or "",
            "polisa_sosotojba": r[2] or "",
            "vk_iznos": float(r[3] or 0),
            "vk_naplata": float(r[4] or 0),
            "saldo": float(r[5] or 0),
        }
        for r in rows
    ]


# ------------------------------------------------------------------
# Проверка 2: обновена полиса → провери дали ги има сите фактури изгенерирано.
# Статусот се зема од ПОСЛЕДНАТА понуда (maxpodbroj), не било кој ponuda ред,
# и се исклучуваат раскинати/сторнирани (42) и капитализирани (18) полиси.
# Исклучени и полисите со префикс "19/" (не подлежат на оваа проверка) —
# иста логика како проверка 3.
# ------------------------------------------------------------------
def check_obnovena_nedostasuvaat_fakturi(godina):
    sql = """
        SELECT p.polisa_broj_cel, a.os_aneks, vrati_godina(a.par_yearid), a.br_rati,
               COUNT(f.os_aneks_fakturaid)
        FROM os_polisa p
        JOIN os_ponuda o ON o.os_ponudaid = p.os_ponudaid
        JOIN os_aneks a ON a.os_polisaid = p.os_polisaid
        LEFT JOIN os_aneks_faktura f ON f.os_aneksid = a.os_aneksid
        WHERE p.status_polisa = 'K'
          AND o.ponuda_podbroj = maxpodbroj_datum(o.ponuda_broj, p.polisa_broj, o.os_produktid, TODAY)
          AND o.par_statusid NOT IN (42, 18)
          AND p.polisa_broj_cel NOT LIKE '19/%'
          AND vrati_godina(a.par_yearid) = ?
        GROUP BY 1, 2, 3, 4
        HAVING COUNT(f.os_aneks_fakturaid) < a.br_rati
    """
    rows = _run(sql, [godina])
    return [
        {"polisa_broj": r[0], "aneks": r[1], "godina": r[2], "br_rati": r[3], "br_generirani": r[4]}
        for r in rows
    ]


# ------------------------------------------------------------------
# Проверка 3: активна полиса → провери дали се изгенерирани сите фактури
# за анексот (без разлика дали се веќе доспеани — фактурите обично се
# генерираат однапред за целата година, па data_faktura < TODAY погрешно
# ги исклучувал веќе-генерираните идни рати и лажно ги пријавувал како
# "недостасуваат").
# Не важи за полиси со еднократна премија (os_produkt_uplata.par_nacin_platiid = 75).
# Исклучени и полисите со префикс "19/" (не подлежат на оваа проверка).
# ------------------------------------------------------------------
def check_aktivna_nedostasuvaat_fakturi(godina):
    sql = """
        SELECT p.polisa_broj_cel, a.os_aneks, vrati_godina(a.par_yearid), a.br_rati,
               COUNT(f.os_aneks_fakturaid)
        FROM os_polisa p
        JOIN os_ponuda o ON o.os_ponudaid = p.os_ponudaid
        JOIN os_produkt_uplata pu ON pu.os_produkt_uplataid = o.os_produkt_uplataid
        JOIN os_aneks a  ON a.os_polisaid = p.os_polisaid
        LEFT JOIN os_aneks_faktura f
               ON f.os_aneksid = a.os_aneksid
        WHERE p.status_polisa = 'K'
          AND o.ponuda_podbroj = maxpodbroj_datum(o.ponuda_broj, p.polisa_broj, o.os_produktid, TODAY)
          AND o.par_statusid IN (17,18)
          AND pu.par_nacin_platiid <> 75
          AND p.polisa_broj_cel NOT LIKE '19/%'
          AND vrati_godina(a.par_yearid) = ?
        GROUP BY 1, 2, 3, 4
        HAVING COUNT(f.os_aneks_fakturaid) < a.br_rati
    """
    rows = _run(sql, [godina])
    return [
        {"polisa_broj": r[0], "aneks": r[1], "godina": r[2], "br_rati": r[3], "br_generirani": r[4]}
        for r in rows
    ]


# ------------------------------------------------------------------
# Проверка 4: фактури со наплата поголема од износот на фактурата.
# Наплатата се зема преку vrati_naplata() UDF, не преку рачен fin_stavka join.
# Собрано на НИВО НА ФАКТУРА (aneks + година + рата), не по одделен
# os_aneks_fakturaid ред — иста логика на групирање како _FAKTURA_EXPR во
# report_fakturi.py. Истата "фактура" (истиот aneks/година-рата) може да
# има повеќе os_aneks_faktura записи (оригинал + сторно/корекција), па
# споредбата per-ред лажно пријавува нарушување иако вкупниот салдо за
# таа фактура е во ред.
#
# Филтрирана по година (a.par_yearid) — без ова vrati_naplata() UDF-от се
# повикуваше по фактура низ ЦЕЛАТА историја на базата (без WHERE опсег),
# што предизвикуваше 504 timeout. Годината доаѓа од истото поле што веќе
# се внесува за проверки 2 и 3.
# ------------------------------------------------------------------
def check_naplata_pogolema_od_faktura(godina):
    sql = """
        SELECT FIRST 1000
               p.polisa_broj_cel,
               MAX(c.desc) AS client_naziv,
               UPPER((a.os_aneks || '/' || TRIM(vrati_godina(a.par_yearid)) || '-' || f.rata)) AS faktura,
               MAX(f.data_faktura) AS data_faktura,
               SUM(NVL(f.iznos,0)) AS iznos,
               SUM(NVL(vrati_naplata(f.os_aneks_fakturaid),0)) AS naplata
        FROM os_aneks_faktura f
        JOIN os_aneks a   ON a.os_aneksid  = f.os_aneksid
        JOIN os_polisa p  ON p.os_polisaid = a.os_polisaid
        JOIN par_client c ON c.par_clientid = a.par_clientid
        WHERE vrati_godina(a.par_yearid) = ?
        GROUP BY 1, 3
        HAVING SUM(NVL(vrati_naplata(f.os_aneks_fakturaid),0)) > SUM(NVL(f.iznos,0))
           AND SUM(NVL(f.iznos,0)) > 0
        ORDER BY 1, 4
    """
    rows = _run(sql, [godina])
    return [
        {
            "polisa_broj": r[0], "client_naziv": r[1], "faktura": r[2],
            "data_faktura": str(r[3]) if r[3] is not None else "",
            "iznos": float(r[4] or 0), "naplata": float(r[5] or 0),
        }
        for r in rows
    ]


# Список на достапни проверки: (клуч, лабела, функција).
# Секоја проверка се извршува одделно и по избор на корисникот — не сите
# наеднаш — за да не се предизвика 504 timeout од долги Informix query-и
# извршени во иста HTTP заявка.
AVAILABLE_CHECKS = [
    ("kapitalizirana_faktura_nula", "1) Статус 18/19/20/36/37/42 — со салдо (иста логика како SMS)", lambda godina: check_kapitalizirana_faktura_nula()),
    ("obnovena_nedostasuvaat", "Полиси — недостасуваат фактури", lambda godina: check_obnovena_nedostasuvaat_fakturi(godina)),
    ("aktivna_nedostasuvaat", "2) Активни — недостасуваат фактури", lambda godina: check_aktivna_nedostasuvaat_fakturi(godina)),
    ("naplata_pogolema_od_faktura", "3) Наплата > износ", lambda godina: check_naplata_pogolema_od_faktura(godina)),
]

_CHECK_KEYS = [key for key, _, _ in AVAILABLE_CHECKS]


def _render(request, user, godina, results, selected_checks, error):
    return templates.TemplateResponse("kontrola_polisi.html", {
        "request": request,
        "user": user,
        "active": "kontrola_polisi",
        "godina": godina,
        "results": results,
        "error": error,
        "searched": results is not None,
        "available_checks": AVAILABLE_CHECKS,
        "selected_checks": selected_checks or [],
    })


@router.get("/kontrola-polisi")
async def kontrola_polisi_get(request: Request):
    user, redirect = _require_role(request)
    if redirect:
        return redirect
    return _render(request, user, date.today().year, None, [], None)


@router.post("/kontrola-polisi")
def kontrola_polisi_post(
    request: Request,
    godina: int = Form(...),
    checks: list[str] = Form(default=[]),
):
    user, redirect = _require_role(request)
    if redirect:
        return redirect

    selected = [key for key in checks if key in _CHECK_KEYS]
    if not selected:
        return _render(request, user, godina, None, [], "Избери барем една проверка.")

    results = {}
    for key, _label, fn in AVAILABLE_CHECKS:
        if key not in selected:
            continue
        try:
            results[key] = fn(godina)
        except Exception as e:
            print(f"[KONTROLA_POLISI] {key} failed: {e}")
            results[key] = {"error": str(e)}

    return _render(request, user, godina, results, selected, None)


# ------------------------------------------------------------------
# Excel извоз за проверка 1 (статус 18/19/20/36/37/42 — со салдо)
# ------------------------------------------------------------------
_HDR_FILL = PatternFill("solid", fgColor="1A3A5C")
_HDR_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
_NORM     = Font(name="Calibri", size=10)
_CTR      = Alignment(horizontal="center", vertical="center")
_RGT      = Alignment(horizontal="right",  vertical="center")
_LFT      = Alignment(horizontal="left",   vertical="center")
_THIN     = Side(border_style="thin", color="CCCCCC")
_BRD      = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)
_NUM      = '#,##0.00'


def _excel_status_saldo(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Статус - салдо"

    headers = ["Полиса", "Статус", "Состојба", "Вк. Износ", "Вк. Наплата", "Салдо"]
    ws.append(headers)
    for ci in range(1, len(headers) + 1):
        c = ws.cell(1, ci)
        c.fill = _HDR_FILL; c.font = _HDR_FONT; c.alignment = _CTR; c.border = _BRD

    for r in rows:
        ws.append([r["polisa_broj"], r["polisa_status_desc"], r["polisa_sosotojba"],
                   r["vk_iznos"], r["vk_naplata"], r["saldo"]])
        dr = ws.max_row
        for ci in (1, 2, 3):
            c = ws.cell(dr, ci); c.font = _NORM; c.border = _BRD; c.alignment = _LFT
        for ci in (4, 5, 6):
            c = ws.cell(dr, ci); c.font = _NORM; c.border = _BRD
            c.alignment = _RGT; c.number_format = _NUM

    for ci, w in [(1, 22), (2, 25), (3, 25), (4, 15), (5, 15), (6, 15)]:
        ws.column_dimensions[get_column_letter(ci)].width = w

    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    return buf


@router.get("/kontrola-polisi/export/status-saldo")
def kontrola_polisi_export_status_saldo(request: Request):
    user, redirect = _require_role(request)
    if redirect:
        return redirect

    rows = check_kapitalizirana_faktura_nula()
    buf = _excel_status_saldo(rows)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Status_Saldo_Dospeana.xlsx"'},
    )
