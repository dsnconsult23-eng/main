from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
import io

from openpyxl.styles import numbers

def format_provision_excel(
    file_path,
    sheet_name="Sheet1",
    prov_col="T",
    label_col="S",
    kurs=0.0,              # курс од vrati_kurs()
    kurs_cell="X1"         # скриена ќелија
):
    wb = load_workbook(file_path)
    ws = wb[sheet_name]

    # 1. Format header
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # 2. Auto width for all columns
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2

    # 3. Fixed widths for specific columns
    fixed_widths = {
        "A": 20,   # Broker
        "D": 20,   # Договорувач
        "E": 15,   # ЕДБ
        "F": 30,   # Адреса
        "G": 15,   # Град
        "H": 25    # Осигуреник
    }
    for col, width in fixed_widths.items():
        ws.column_dimensions[col].width = width

    # 4. Add total provision row
    last_row = ws.max_row + 1
     # ------------------------------------------------
    # 1️⃣ Стави курс во скриена ќелија
    # ------------------------------------------------
    ws[kurs_cell] = float(kurs)
    ws.column_dimensions["X"].hidden = True   # цела колона скриена

    # ------------------------------------------------
    # 2️⃣ Додај САМО ЕДЕН ред на крај
    # ------------------------------------------------
    last_data_row = ws.max_row

    total_row = last_data_row + 1
    ws.insert_rows(total_row)

    ws[f"{label_col}{total_row}"] = "Вкупно провизија (МКД):"
    ws[f"{label_col}{total_row}"].font = Font(bold=True)

    # FORMULA (НЕ директно множење со број!)
    ws[f"{prov_col}{total_row}"] = (
        f"=SUM({prov_col}2:{prov_col}{last_data_row})*{kurs_cell}"
    )
    ws[f"{prov_col}{total_row}"].font = Font(bold=True)
    ws[f"{prov_col}{total_row}"].number_format = "#,##0.00"
    # Border
    thin = Side(style="thin", color="000000")
    for c in range(1, ws.max_column + 1):
        ws.cell(row=last_row, column=c).border = Border(top=thin, bottom=thin)

    wb.save(file_path)



def format_provision_excel_in_memory(excel_io: io.BytesIO, sheet_name="Sheet1", prov_col="T" , 
    label_col="S",
    kurs=0.0,              # курс од vrati_kurs()
    kurs_cell="X1"         # скриена ќелија
    ):
    """
    Форматира Excel од BytesIO и враќа нов BytesIO објект.
    """
    excel_io.seek(0)
    wb = load_workbook(excel_io)
    ws = wb[sheet_name]

    # 1. Format header
    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # 2. Auto width for all columns
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2

    # 3. Fixed widths for специфични колони
    fixed_widths = {
        "A": 20,   # Broker
        "D": 20,   # Договорувач
        "E": 15,   # ЕДБ
        "F": 30,   # Адреса
        "G": 15,   # Град
        "H": 25    # Осигуреник
    }
    for col, width in fixed_widths.items():
        ws.column_dimensions[col].width = width

     # ------------------------------------------------
    # 1️⃣ Стави курс во скриена ќелија
    # ------------------------------------------------
    last_data_row = ws.max_row
    ws[kurs_cell] = float(kurs)
    ws.column_dimensions["X"].hidden = True   # цела колона скриена

    # ------------------------------------------------
    # 2️⃣ Додај САМО ЕДЕН ред на крај
    # ------------------------------------------------
    total_row = last_data_row + 1
    ws.insert_rows(total_row)

    ws[f"{label_col}{total_row}"] = "Вкупно провизија (МКД):"
    ws[f"{label_col}{total_row}"].font = Font(bold=True)

    # FORMULA (НЕ директно множење со број!)
    ws[f"{prov_col}{total_row}"] = (
        f"=SUM({prov_col}2:{prov_col}{last_data_row})*{kurs_cell}"
    )
    ws[f"{prov_col}{total_row}"].font = Font(bold=True)
    ws[f"{prov_col}{total_row}"].number_format = "#,##0.00"

    # Border
    thin = Side(style="thin", color="000000")
    for c in range(1, ws.max_column + 1):
        ws.cell(row=last_data_row, column=c).border = Border(top=thin, bottom=thin)

    # Save back to BytesIO
    out_io = io.BytesIO()
    wb.save(out_io)
    out_io.seek(0)
    return out_io

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def format_full_excel_sms(file_path, sheet_name="Sheet1"):
    wb = load_workbook(file_path)
    ws = wb[sheet_name]

    # ----------------------------------------
    # 1. HEADER STYLE
    # ----------------------------------------
    header_font = Font(bold=True)
    thin = Side(border_style="thin", color="000000")
    header_border = Border(top=thin, left=thin, right=thin, bottom=thin)

    # Боја на header редот (жолто)
    header_fill = PatternFill(start_color="00FFFF00", end_color="00FFFF00", fill_type="solid")

    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.fill = header_fill
        cell.border = header_border

    # ----------------------------------------
    # 2. FORMAT FOR SPECIFIC COLUMNS
    # ----------------------------------------
    date_columns = ["C", "Q"]  # пример: Datum pocetka, Premija placena do
    money_columns = ["L", "M", "N", "O"]  # пример: износи

    for col in date_columns:
        for cell in ws[col][1:]:
            cell.number_format = "DD.MM.YYYY"

    for col in money_columns:
        for cell in ws[col][1:]:
            cell.number_format = "#.##0,00"

    # ----------------------------------------
    # 3. AUTO-WIDTH
    # ----------------------------------------
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2

    # ----------------------------------------
    # 4. FREEZE HEADER + AUTO FILTER
    # ----------------------------------------
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ----------------------------------------
    # 5. BORDER FOR ALL CELLS
    # ----------------------------------------
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = cell_border

    # ----------------------------------------
    # 6. SAVE
    # ----------------------------------------
    wb.save(file_path)
    print("✔ Excel е успешно форматиран:", file_path)

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# ---------------------------------------------------------
# LATINICA MAPPING (bez dijakritika)
# ---------------------------------------------------------
MK_TO_LAT = {
    "А":"A","а":"a",
    "Б":"B","б":"b",
    "В":"V","в":"v",
    "Г":"G","г":"g",
    "Д":"D","д":"d",
    "Ѓ":"Gj","ѓ":"gj",
    "Е":"E","е":"e",
    "Ж":"Z","ж":"z",
    "З":"Z","з":"z",
    "Ѕ":"Dz","ѕ":"dz",
    "И":"I","и":"i",
    "Ј":"J","ј":"j",
    "К":"K","к":"k",
    "Л":"L","л":"l",
    "Љ":"Lj","љ":"lj",
    "М":"M","м":"m",
    "Н":"N","н":"n",
    "Њ":"Nj","њ":"nj",
    "О":"O","о":"o",
    "П":"P","п":"p",
    "Р":"R","р":"r",
    "С":"S","с":"s",
    "Т":"T","т":"t",
    "Ќ":"Kj","ќ":"kj",
    "У":"U","у":"u",
    "Ф":"F","ф":"f",
    "Х":"H","х":"h",
    "Ц":"C","ц":"c",
    "Ч":"C","ч":"c",
    "Џ":"Dj","џ":"dj",
    "Ш":"S","ш":"s"
}

def mk_to_latin(text):
    if not isinstance(text, str):
        return text
    res = ""
    for ch in text:
        res += MK_TO_LAT.get(ch, ch)
    return res


# ---------------------------------------------------------
# MAIN FORMATTER
# ---------------------------------------------------------
def format_broker_excel(file_path, sheet_name="Sheet1"):
    wb = load_workbook(file_path)
    ws = wb[sheet_name]

    thin = Side(border_style="thin", color="000000")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)

    # ----------------------------
    # 1. HEADER FORMAT
    # ----------------------------
    yellow_cols = ["T", "U", "V", "W", "Y", "Z"]
    red_cols = ["Q"]

    yellow_fill = PatternFill(start_color="00FFFF00", end_color="00FFFF00", fill_type="solid")
    red_fill = PatternFill(start_color="00FF0000", end_color="00FF0000", fill_type="solid")

    for cell in ws[1]:
        letter = cell.column_letter
        cell.font = bold
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_all

        if letter in yellow_cols:
            cell.fill = yellow_fill
        elif letter in red_cols:
            cell.fill = red_fill

    # ----------------------------
    # 2. DATE FORMAT
    # ----------------------------
    date_cols = ["U", "V", "AB", "AC", "C", "Q"]
    for col in date_cols:
        for cell in ws[col][1:]:
            cell.number_format = "DD.MM.YYYY"

    # ----------------------------
    # 3. MONEY FORMAT


    money_cols = ["T", "W"]

    for col in money_cols:
        for cell in ws[col][1:]:  # skip header
            raw = cell.value

            if raw is None:
                continue

            # ▶ STEP 1: convert text → float
            if isinstance(raw, str):
                clean = raw.strip()
                clean = clean.replace(" ", "")
                clean = clean.replace(".", "").replace(",", ".")

                try:
                    raw = float(clean)
                except:
                    continue

            # ▶ STEP 2: force rounding to EXACT 2 decimals
            raw = round(float(raw), 2)
            cell.value = raw

            # ▶ STEP 3: Excel european number format (препорачано)
            cell.number_format = "#,##0.00"




    # ----------------------------
    # 4. AUTO WIDTH
    # ----------------------------
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2
    # ----------------------------------------------------
    # 4a. FIXED WIDTH overrides
    # ----------------------------------------------------
    ws.column_dimensions["A"].width = 12   # пример вредност – можеш да смениш
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 25
    ws.column_dimensions["H"].width = 25
    ws.column_dimensions["K"].width = 18
    ws.column_dimensions["AH"].width = 18

    # ----------------------------
    # 5. FREEZE + FILTER
    # ----------------------------
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ----------------------------
    # 6. BORDER FOR ALL DATA ROWS
    # ----------------------------
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border_all

    # ----------------------------
    # 7. TOTAL SUM IN COLUMN T
    # ----------------------------
    last_row = ws.max_row + 1

    ws[f"S{last_row}"] = "Vkupno provizija:"
    ws[f"S{last_row}"].font = bold

    ws[f"T{last_row}"] = f"=SUM(T2:T{last_row-1})"
    ws[f"T{last_row}"].font = bold
    ws[f"T{last_row}"].number_format = "#.##0,00"

    for col in range(1, ws.max_column + 1):
        ws.cell(row=last_row, column=col).border = border_all

    wb.save(file_path)
    print("✔ Excel formatiran.")


# ---------------------------------------------------------
# 3) CONVERT FULL EXCEL TO LATINICA
# ---------------------------------------------------------
def convert_excel_to_latin(file_path, sheet_name="Sheet1"):
    wb = load_workbook(file_path)
    ws = wb[sheet_name]

    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.value = mk_to_latin(cell.value)

    wb.save(file_path)
    print("✔ Teskstot e konvertiran vo latinica.")


# ---------------------------------------------------------
# 4) MASTER FUNCTION – CALL ONE FUNCTION
# ---------------------------------------------------------
def process_broker_excel(file_path):
    convert_excel_to_latin(file_path)
    format_broker_excel(file_path)
    print("✔ Zavrseno kompletno.")

import io
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime

def format_broker_excel_in_memory(excel_io: io.BytesIO, sheet_name="Sheet1", prov_col="T"):
    """
    НОВА ФУНКЦИЈА:
    Форматира Excel од BytesIO и враќа нов BytesIO со комплет форматирање.
    """

    excel_io.seek(0)
    wb = load_workbook(excel_io)
    ws = wb[sheet_name]

    thin = Side(style="thin", color="000000")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)

    # ----------------------------------------------------
    # 1. HEADER FORMATTING
    # ----------------------------------------------------
    yellow_fill = PatternFill(start_color="00FFFF00", end_color="00FFFF00", fill_type="solid")
    red_fill    = PatternFill(start_color="00FF0000", end_color="00FF0000", fill_type="solid")

    yellow_cols = ["T", "U", "V", "W", "Y", "Z"]
    red_cols = ["Q"]

    for cell in ws[1]:
        col = cell.column_letter
        cell.font = bold
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_all

        if col in yellow_cols:
            cell.fill = yellow_fill
        if col in red_cols:
            cell.fill = red_fill

    # ----------------------------------------------------
    # 2. DATE FORMAT FOR SPECIFIC COLUMNS
    # ----------------------------------------------------
    date_cols = ["U", "V", "AB", "AC", "C", "Q"]

    for col in date_cols:
        for cell in ws[col][1:]:
            val = cell.value
            if isinstance(val, str):
                val = val.strip()
                try:
                    cell.value = datetime.strptime(val, "%d/%m/%Y")
                except:
                    try:
                        cell.value = datetime.strptime(val, "%Y-%m-%d")
                    except:
                        pass
            cell.number_format = "DD.MM.YYYY"

    # ----------------------------------------------------
    # 3. MONEY FORMATTING (T, W)
    # ----------------------------------------------------
    money_cols = ["T", "W"]

    for col in money_cols:
        for cell in ws[col][1:]:
            v = cell.value
            if isinstance(v, str):
                clean = v.replace(".", "").replace(",", ".")
                try:
                    v = float(clean)
                except:
                    continue

            if isinstance(v, (float, int)):
                v = round(v, 2)
                cell.value = v

            cell.number_format = "#,##0.00"

    # ----------------------------------------------------
    # 4. AUTO WIDTH
    # ----------------------------------------------------
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_len + 2

    # ----------------------------------------------------
    # 4a. FIXED WIDTH FOR F, G, H, K, AH
    # ----------------------------------------------------
    fixed_cols = {
        "F": 20,
        "G": 25,
        "H": 25,
        "K": 18,
        "AH": 18
    }

    for col, width in fixed_cols.items():
        ws.column_dimensions[col].width = width

    # ----------------------------------------------------
    # 5. FREEZE + FILTER
    # ----------------------------------------------------
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ----------------------------------------------------
    # 6. BORDERS FOR ALL ROWS
    # ----------------------------------------------------
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border_all

    # ----------------------------------------------------
    # 7. TOTAL SUM IN T
    # ----------------------------------------------------
    last_row = ws.max_row + 1

    ws[f"S{last_row}"] = "Vkupno provizija:"
    ws[f"S{last_row}"].font = bold

    ws[f"T{last_row}"] = f"=SUM(T2:T{last_row-1})"
    ws[f"T{last_row}"].font = bold
    ws[f"T{last_row}"].number_format = "#,##0.00"

    for col in range(1, ws.max_column + 1):
        ws.cell(row=last_row, column=col).border = border_all

    # ----------------------------------------------------
    # 8. RETURN NEW BYTES
    # ----------------------------------------------------
    out_io = io.BytesIO()
    wb.save(out_io)
    out_io.seek(0)
    return out_io






